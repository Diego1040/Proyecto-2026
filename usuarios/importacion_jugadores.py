"""Validación de solo lectura e importación explícita de jugadores históricos."""

import math
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time
from numbers import Real
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.datetime import CALENDAR_WINDOWS_1900, from_excel
from openpyxl.utils.exceptions import InvalidFileException

from .models import Apoderado, Jugador, Usuario
from .services import asignar_categoria_automatica, obtener_categoria_automatica
from .validators import normalizar_rut, validar_rut


ENCABEZADOS = (
    "N°", "NOMBRES", "APELLIDOS", "F.N", "CAT.", "CONTACTO", "RUT",
    "AVTIVO", "GENERO", "FEDERADO",
)
ABREVIATURAS = ("U7", "U9", "U11", "U13", "U15", "U17", "T/C")
RAMAS = {"F": Jugador.Rama.DAMAS, "M": Jugador.Rama.VARONES}
ESTADOS = {"SI": Jugador.Estado.ACTIVO, "NO": Jugador.Estado.INACTIVO}


class ErrorArchivoJugadores(ValueError):
    """Estructura o contenido XLSX que no permite examinar las filas."""


class ErrorImportacionJugadores(Exception):
    def __init__(self, mensaje, informe):
        super().__init__(mensaje)
        self.informe = informe


@dataclass(frozen=True)
class PoliticaImportacion:
    """Decisiones centralizadas; None significa pendiente, nunca un default."""

    estado_vacio: str | None = Jugador.Estado.INACTIVO
    procedencia: str | None = ""
    federado_vacio: bool = False


@dataclass(frozen=True)
class Incidencia:
    codigo: str
    mensaje: str
    nivel: str = "INFO"


@dataclass
class ResultadoFila:
    fila: int
    numero: object
    originales: dict
    normalizados: dict = field(default_factory=dict)
    categoria_calculada: str | None = None
    incidencias: list[Incidencia] = field(default_factory=list)
    existente: bool = False
    conflicto: bool = False
    omitida: bool = False
    categoria_id: int | None = None
    jugador_id: int | None = None
    creada: bool = False

    def agregar(self, codigo, mensaje, nivel="INFO"):
        self.incidencias.append(Incidencia(codigo, mensaje, nivel))

    @property
    def estado(self):
        if self.omitida:
            return "OMITIDA"
        if self.creada:
            return "CREADO"
        if any(i.nivel == "REVISION" for i in self.incidencias):
            return "REVISIÓN"
        if any(i.nivel == "PENDIENTE" for i in self.incidencias):
            return "POLÍTICA PENDIENTE"
        if self.existente:
            return "EXISTENTE SIN CAMBIOS"
        return "OK"


@dataclass
class InformeJugadores:
    temporada: int
    fecha_referencia: date
    encabezados: tuple
    filas: list[ResultadoFila]
    modo: str = "validar"
    rollback: bool = False
    seleccionadas: int = 0

    @property
    def resumen(self):
        # Existentes/conflictos son indicadores que pueden coincidir con revisión.
        return {
            "filas examinadas": len(self.filas),
            "jugadores detectados": sum(not f.omitida for f in self.filas),
            "OK": sum(f.estado == "OK" for f in self.filas),
            "importables": self.seleccionadas if self.modo == "importar" else sum(f.estado == "OK" for f in self.filas),
            "creados": sum(f.creada for f in self.filas),
            "revisión": sum(f.estado == "REVISIÓN" for f in self.filas),
            "política pendiente": sum(f.estado == "POLÍTICA PENDIENTE" for f in self.filas),
            "existentes": sum(f.existente for f in self.filas),
            "existentes sin cambios": sum(f.estado == "EXISTENTE SIN CAMBIOS" for f in self.filas),
            "conflictos": sum(f.conflicto for f in self.filas),
            "omitidas": sum(f.omitida for f in self.filas),
        }


def _texto(valor):
    return "" if valor is None else str(valor).strip()


def _encabezados(celdas):
    originales = tuple(c.value for c in celdas)
    posiciones = {}
    for posicion, celda in enumerate(celdas):
        nombre = _texto(celda.value).upper()
        if not nombre:
            continue
        if celda.data_type == "f":
            raise ErrorArchivoJugadores("Los encabezados no pueden contener fórmulas.")
        if nombre in posiciones:
            raise ErrorArchivoJugadores(f"Encabezado duplicado: {nombre}.")
        if nombre not in ENCABEZADOS:
            raise ErrorArchivoJugadores(f"Encabezado no reconocido: {nombre}.")
        posiciones[nombre] = posicion
    faltantes = [nombre for nombre in ENCABEZADOS if nombre not in posiciones]
    if faltantes:
        raise ErrorArchivoJugadores("Encabezados ausentes: " + ", ".join(faltantes) + ".")
    return originales, posiciones


def _leer_xlsx(archivo):
    libro = None
    try:
        libro = load_workbook(archivo, read_only=True, data_only=False, keep_links=False)
        if "Hoja1" not in libro.sheetnames:
            raise ErrorArchivoJugadores("No existe la hoja requerida Hoja1.")
        hoja = libro["Hoja1"]
        # No confiar en dimensiones o tablas internas: el XLSX del club es disperso.
        hoja.reset_dimensions()
        iterador = hoja.iter_rows()
        cabecera = next(iterador, ())
        originales, posiciones = _encabezados(cabecera)
        filas = []
        for numero_fila, celdas in enumerate(iterador, start=2):
            valores, formulas, errores = {}, set(), set()
            for nombre, posicion in posiciones.items():
                celda = celdas[posicion] if posicion < len(celdas) else None
                valores[nombre] = celda.value if celda is not None else None
                if celda is not None and celda.data_type == "f":
                    formulas.add(nombre)
                if celda is not None and celda.data_type == "e":
                    errores.add(nombre)
            filas.append((numero_fila, valores, formulas, errores))
        return originales, libro.epoch, filas
    except (OSError, BadZipFile, InvalidFileException, ParseError, KeyError, ValueError, TypeError) as error:
        if isinstance(error, ErrorArchivoJugadores):
            raise
        raise ErrorArchivoJugadores(f"No se pudo leer el archivo XLSX: {error}") from error
    finally:
        if libro is not None:
            libro.close()


def normalizar_fecha_excel(valor, epoch):
    """Convierte sólo representaciones explícitas; nunca adivina año u origen."""
    if isinstance(valor, datetime):
        if valor.time() != time.min:
            raise ValueError("F.N contiene una hora distinta de medianoche.")
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, Real) and not isinstance(valor, bool):
        if not math.isfinite(valor) or valor <= 0:
            raise ValueError("Serial Excel de F.N inválido.")
        if epoch == CALENDAR_WINDOWS_1900 and int(valor) == 60:
            raise ValueError("El serial 60 representa una fecha inexistente en Excel.")
        convertido = from_excel(valor, epoch=epoch)
        if not isinstance(convertido, datetime):
            raise ValueError("F.N no representa una fecha completa.")
        return normalizar_fecha_excel(convertido, epoch)
    if isinstance(valor, str):
        texto = valor.strip()
        for patron, formato in (
            (r"[0-9]{4}-[0-9]{2}-[0-9]{2}", "%Y-%m-%d"),
            (r"[0-9]{2}/[0-9]{2}/[0-9]{4}", "%d/%m/%Y"),
            (r"[0-9]{2}-[0-9]{2}-[0-9]{4}", "%d-%m-%Y"),
        ):
            if re.fullmatch(patron, texto):
                return datetime.strptime(texto, formato).date()
    raise ValueError("F.N inválida o ambigua; use una fecha Excel o texto con año completo.")


def _referencia_categoria(valor):
    texto = " ".join(_texto(valor).upper().split())
    if texto in ABREVIATURAS:
        return texto, None
    coincidencia = re.fullmatch(
        r"(U7|U9|U11|U13|U15|U17|T/C|TODO COMPETIDOR)(?: (DAMAS|VARONES|MIXTO))?",
        texto,
    )
    if coincidencia:
        familia, rama = coincidencia.groups()
        return ("T/C" if familia == "TODO COMPETIDOR" else familia), rama
    return None


def _indice_existentes(modelo):
    indice = defaultdict(list)
    for objeto in modelo.objects.all():
        indice[normalizar_rut(objeto.rut)].append(objeto)
    return indice


def _normalizar_fila(resultado, valores, politica, epoch, fecha_referencia):
    datos = resultado.normalizados
    for columna, campo in (("NOMBRES", "nombres"), ("APELLIDOS", "apellidos")):
        valor = valores[columna]
        datos[campo] = _texto(valor) if isinstance(valor, str) else ""
        if not _texto(valor):
            resultado.agregar(campo.upper() + "_AUSENTE", f"{columna} ausente.", "REVISION")
        elif not isinstance(valor, str):
            resultado.agregar(campo.upper() + "_INVALIDO", f"{columna} debe ser texto.", "REVISION")
        else:
            try:
                Jugador._meta.get_field(campo).clean(datos[campo], None)
            except ValidationError as error:
                resultado.agregar(campo.upper() + "_INVALIDO", f"{columna}: {'; '.join(error.messages)}", "REVISION")

    datos["fecha_nacimiento"] = None
    if not _texto(valores["F.N"]):
        resultado.agregar("FECHA_AUSENTE", "F.N ausente.", "REVISION")
    else:
        try:
            nacimiento = normalizar_fecha_excel(valores["F.N"], epoch)
            if nacimiento > fecha_referencia:
                raise ValueError("La fecha de nacimiento no puede estar en el futuro.")
            datos["fecha_nacimiento"] = nacimiento
        except (ValueError, TypeError, OverflowError) as error:
            resultado.agregar("FECHA_INVALIDA", str(error), "REVISION")

    datos["rut"] = ""
    if not _texto(valores["RUT"]):
        resultado.agregar("RUT_AUSENTE", "RUT ausente.", "REVISION")
    else:
        try:
            valor = valores["RUT"]
            if isinstance(valor, bool) or not isinstance(valor, (str, int, float)):
                raise ValueError("RUT debe ser texto o un número entero.")
            if isinstance(valor, float):
                if not math.isfinite(valor) or not valor.is_integer():
                    raise ValueError("RUT numérico inválido.")
                valor = int(valor)
            rut = normalizar_rut(valor)
            Jugador._meta.get_field("rut").clean(rut, None)
            validar_rut(rut)
            datos["rut"] = rut
        except (ValidationError, ValueError, TypeError) as error:
            mensajes = error.messages if isinstance(error, ValidationError) else [str(error)]
            resultado.agregar("RUT_INVALIDO", "; ".join(mensajes), "REVISION")

    genero = _texto(valores["GENERO"]).upper()
    datos["rama"] = RAMAS.get(genero)
    if datos["rama"] is None:
        resultado.agregar("GENERO_INVALIDO", f"GENERO {genero or 'vacío'}: se requiere M o F.", "REVISION")

    contacto = valores["CONTACTO"]
    datos["telefono"] = re.sub(r"[\s()\-]", "", _texto(contacto))
    if _texto(contacto):
        if not isinstance(contacto, str) or not re.fullmatch(r"\+569[0-9]{8}", datos["telefono"]):
            datos["telefono"] = ""
            resultado.agregar(
                "CONTACTO_INVALIDO",
                "CONTACTO inválido: no es un teléfono móvil chileno completo (+569 y 8 dígitos). "
                "Administración debe revisar el valor original del reporte; no se corrige ni persiste y no bloquea al jugador.",
                "ADMINISTRATIVA",
            )
        else:
            resultado.agregar("CONTACTO_LEIDO", f"CONTACTO normalizado={datos['telefono']}.")
    else:
        resultado.agregar("CONTACTO_AUSENTE", "CONTACTO vacío; no se inventa.")

    activo = _texto(valores["AVTIVO"]).upper()
    datos["estado"] = ESTADOS.get(activo) if activo else politica.estado_vacio
    if not activo and datos["estado"] is None:
        resultado.agregar("ACTIVO_PENDIENTE", "AVTIVO vacío: política pendiente.", "PENDIENTE")
    elif activo not in ESTADOS and activo:
        resultado.agregar("ACTIVO_INVALIDO", f"AVTIVO={activo}: valor no reconocido.", "REVISION")
    else:
        resultado.agregar("ACTIVO_LEIDO", f"AVTIVO={activo or 'vacío'} -> {datos['estado']}.")

    federado = _texto(valores["FEDERADO"]).upper()
    datos["federado"] = {"SI": True, "NO": False, "": politica.federado_vacio}.get(federado)
    if federado not in ("", "SI", "NO"):
        resultado.agregar("FEDERADO_INVALIDO", f"FEDERADO={federado}: valor no reconocido.", "REVISION")
    else:
        resultado.agregar("FEDERADO_LEIDO", f"FEDERADO={federado or 'vacío'} -> {datos['federado']}; solo reporte, no se persiste.")

    datos["procedencia"] = politica.procedencia
    if politica.procedencia is None:
        resultado.agregar("PROCEDENCIA_PENDIENTE", "Procedencia ausente: política pendiente; no se asume INDEPENDIENTE.", "PENDIENTE")


def _calcular_categoria(resultado, valores, temporada, fecha_referencia):
    datos = resultado.normalizados
    referencia = _referencia_categoria(valores["CAT."])
    if not _texto(valores["CAT."]):
        resultado.agregar("CATEGORIA_AUSENTE", "CAT. vacía: no se puede comparar la referencia.")
    elif referencia is None:
        resultado.agregar("CATEGORIA_DESCONOCIDA", f"CAT.={_texto(valores['CAT.'])}: referencia no reconocida.", "REVISION")
    categoria = None
    if datos["fecha_nacimiento"] is not None and datos["rama"] is not None:
        try:
            categoria = obtener_categoria_automatica(
                datos["fecha_nacimiento"], datos["rama"],
                temporada=temporada, fecha_referencia=fecha_referencia,
            )
            if categoria is None:
                resultado.agregar("SIN_REGLA_CATEGORIA", f"No hay categoría aplicable para la temporada {temporada}.", "REVISION")
            else:
                resultado.categoria_calculada = categoria.nombre
                resultado.categoria_id = categoria.pk
                calculada = _referencia_categoria(categoria.nombre)
                if referencia and (
                    calculada is None or referencia[0] != calculada[0]
                    or (referencia[1] is not None and referencia[1] != categoria.rama)
                ):
                    resultado.agregar("CATEGORIA_DISCREPANTE", f"Categoría Excel {_texto(valores['CAT.'])} / calculada {categoria.nombre}; requiere revisión, sin excepción automática.", "REVISION")
        except ValidationError as error:
            resultado.agregar("REGLAS_CATEGORIA_INVALIDAS", "; ".join(error.messages), "REVISION")
    return categoria


def _preparar_contacto(resultado, categoria):
    datos = resultado.normalizados
    datos["telefono_jugador"] = ""
    datos["contacto_destino"] = "SIN DETERMINAR"
    if categoria is None:
        return
    referencia = _referencia_categoria(categoria.nombre)
    # La titularidad depende de la categoría calculada, nunca de CAT. del archivo.
    if referencia and referencia[0] == "T/C":
        datos["telefono_jugador"] = datos["telefono"]
        datos["contacto_destino"] = "JUGADOR"
    else:
        datos["contacto_destino"] = "APODERADO"
        if datos["telefono"]:
            resultado.agregar("CONTACTO_APODERADO", f"Teléfono de apoderado {datos['telefono']}: no persistido. Administración debe completar los datos del apoderado posteriormente.")


def _comparar_existentes(resultado, indices, categoria):
    datos = resultado.normalizados
    for nombre, indice in indices.items():
        objetos = indice.get(datos["rut"], [])
        if not objetos:
            continue
        resultado.existente = True
        jugadores = indices["Jugador"].get(datos["rut"], [])
        if (nombre == "Usuario" and len(objetos) == len(jugadores) == 1
                and jugadores[0].usuario_id == objetos[0].pk):
            resultado.agregar("USUARIO_YA_VINCULADO", "Cuenta ya vinculada al jugador existente; no se modifica el vínculo.")
            continue
        if nombre != "Jugador" or len(objetos) > 1:
            resultado.conflicto = True
            resultado.agregar("RUT_EN_" + nombre.upper(), f"RUT presente en {nombre} (IDs {', '.join(str(o.pk) for o in objetos)}); revisión de identidad, sin vincular ni modificar.", "REVISION")
            continue
        jugador = objetos[0]
        resultado.jugador_id = jugador.pk
        diferencias = []
        for campo in ("nombres", "apellidos", "fecha_nacimiento", "rama"):
            entrante, actual = datos[campo], getattr(jugador, campo)
            if isinstance(entrante, str) and isinstance(actual, str):
                entrante, actual = entrante.casefold(), actual.strip().casefold()
            if entrante and entrante != actual:
                diferencias.append(campo)
        if datos["estado"] and datos["estado"] != jugador.estado:
            diferencias.append("estado")
        telefono_actual = re.sub(r"[\s()\-]", "", jugador.telefono)
        if datos.get("telefono_jugador") and datos["telefono_jugador"] != telefono_actual:
            diferencias.append("telefono")
        if categoria is not None and jugador.categoria_actual_id != categoria.pk:
            diferencias.append("categoria_actual (conservar asignación/excepción existente)")
        if diferencias:
            resultado.conflicto = True
            resultado.agregar("JUGADOR_DIFERENTE", f"RUT ya pertenece a Jugador #{jugador.pk}; diferencias en {', '.join(diferencias)}. No se actualiza.", "REVISION")
        else:
            resultado.agregar("JUGADOR_EXISTENTE", f"Jugador #{jugador.pk} existente; sin diferencias detectadas en los campos comparables. No se modifica.")


def _construir_jugador(resultado, categoria=None):
    datos = resultado.normalizados
    return Jugador(
        nombres=datos["nombres"], apellidos=datos["apellidos"],
        fecha_nacimiento=datos["fecha_nacimiento"], rut=datos["rut"], rama=datos["rama"],
        telefono=datos["telefono_jugador"], categoria_actual=categoria,
        estado=datos["estado"] or "", procedencia=datos["procedencia"] or "",
        club_anterior="", usuario=None, fecha_ingreso=None,
        email="", nombre_contacto_emergencia="", telefono_contacto_emergencia="",
    )


def _validar_modelo_en_memoria(resultado, categoria):
    datos = resultado.normalizados
    if not all(datos[c] for c in ("nombres", "apellidos", "fecha_nacimiento", "rut", "rama")):
        return
    candidato = _construir_jugador(resultado, categoria)
    # Estos dos campos se reportan por separado si están sin confirmar.
    # La unicidad se compara explícitamente, incluyendo RUT históricos no canónicos.
    excluidos = [c for c in ("estado", "procedencia") if datos[c] is None]
    try:
        candidato.full_clean(exclude=excluidos, validate_unique=False)
    except ValidationError as error:
        if hasattr(error, "message_dict"):
            mensaje = "; ".join(f"{campo}: {', '.join(mensajes)}" for campo, mensajes in error.message_dict.items())
        else:
            mensaje = "; ".join(error.messages)
        resultado.agregar("VALIDACION_MODELO", mensaje, "REVISION")


def validar_archivo_jugadores(archivo, *, temporada, politica=None, fecha_referencia=None):
    """Lee el XLSX y consulta la BD. No guarda candidatos ni registros auxiliares."""
    if not isinstance(temporada, int) or isinstance(temporada, bool) or not 1 <= temporada <= 9999:
        raise ErrorArchivoJugadores("Indique una temporada válida entre 1 y 9999.")
    politica = politica or PoliticaImportacion()
    fecha_referencia = fecha_referencia or timezone.localdate()
    encabezados, epoch, leidas = _leer_xlsx(archivo)
    indices = {modelo.__name__: _indice_existentes(modelo) for modelo in (Jugador, Apoderado, Usuario)}
    resultados, ruts = [], defaultdict(list)
    for numero_fila, originales, formulas, errores in leidas:
        resultado = ResultadoFila(numero_fila, originales["N°"], originales)
        resultados.append(resultado)
        if not any(_texto(originales[c]) for c in ENCABEZADOS if c != "N°"):
            resultado.omitida = True
            resultado.agregar("SIN_DATOS_JUGADOR", "Fila sin datos de jugador (vacía o solamente N°).")
            continue
        valores = originales.copy()
        for columna in sorted(formulas | errores):
            if columna == "CONTACTO":
                # Se valida el texto original sin evaluarlo; su invalidez no bloquea.
                continue
            valores[columna] = None
            resultado.agregar("CELDA_NO_LITERAL", f"{columna} contiene fórmula o error Excel; no se evalúa.", "REVISION")
        _normalizar_fila(resultado, valores, politica, epoch, fecha_referencia)
        categoria = _calcular_categoria(resultado, valores, temporada, fecha_referencia)
        _preparar_contacto(resultado, categoria)
        _comparar_existentes(resultado, indices, categoria)
        _validar_modelo_en_memoria(resultado, categoria)
        if resultado.normalizados["rut"]:
            ruts[resultado.normalizados["rut"]].append(resultado)
    for grupo in ruts.values():
        if len(grupo) > 1:
            filas = ", ".join(str(f.fila) for f in grupo)
            for resultado in grupo:
                resultado.conflicto = True
                resultado.agregar("RUT_DUPLICADO", f"RUT duplicado dentro del XLSX en filas {filas}; se revisa todo el grupo.", "REVISION")
    return InformeJugadores(temporada, fecha_referencia, encabezados, resultados)


def importar_archivo_jugadores(archivo, *, temporada, fecha_referencia=None):
    """Revalida siempre. Importa sólo OK; una falla revierte todo el lote elegido."""
    informe = validar_archivo_jugadores(
        archivo, temporada=temporada, fecha_referencia=fecha_referencia,
    )
    informe.modo = "importar"
    seleccionadas = [fila for fila in informe.filas if fila.estado == "OK"]
    informe.seleccionadas = len(seleccionadas)
    if not seleccionadas:
        return informe
    creados = []
    fila_actual = None
    try:
        with transaction.atomic():
            # Repetir coincidencias tras la previsualización. La unicidad de BD
            # y full_clean protegen además las escrituras concurrentes.
            indices = {modelo.__name__: _indice_existentes(modelo) for modelo in (Jugador, Apoderado, Usuario)}
            for fila in seleccionadas:
                fila_actual = fila.fila
                comprobacion = ResultadoFila(fila.fila, fila.numero, fila.originales, fila.normalizados)
                _comparar_existentes(comprobacion, indices, None)
                if comprobacion.existente:
                    raise ValidationError("La identidad cambió en BD después de validar. Vuelva a revisar el archivo.")
                jugador = _construir_jugador(fila)
                jugador.full_clean()
                jugador.save()
                categoria = asignar_categoria_automatica(
                    jugador, temporada=temporada, fecha_referencia=informe.fecha_referencia,
                )
                if categoria is None or categoria.pk != fila.categoria_id:
                    raise ValidationError("La categoría cambió después de validar. Revise las reglas de la temporada.")
                jugador.full_clean()
                creados.append((fila, jugador.pk))
    except Exception as error:
        informe.rollback = True
        mensaje = f"Fallo en fila {fila_actual}; rollback completo del conjunto seleccionado: {error}"
        for fila in seleccionadas:
            fila.agregar("LOTE_REVERTIDO", mensaje, "REVISION")
        raise ErrorImportacionJugadores(mensaje, informe) from error
    # Sólo anunciar creaciones cuando la transacción haya terminado correctamente.
    for fila, jugador_id in creados:
        fila.creada = True
        fila.jugador_id = jugador_id
        fila.agregar("JUGADOR_CREADO", f"Jugador #{jugador_id} creado con historial automático de categoría.")
    return informe


def comprobar_destino_reporte(destino, archivo):
    ruta = Path(destino).resolve()
    if ruta == Path(archivo).resolve():
        raise ErrorArchivoJugadores("El reporte no puede sobrescribir el archivo de entrada.")
    if ruta.suffix.lower() != ".xlsx":
        raise ErrorArchivoJugadores("El reporte debe tener extensión .xlsx.")
    if ruta.exists():
        raise ErrorArchivoJugadores("El reporte ya existe. Elija otra ruta para conservarlo.")
    if not ruta.parent.is_dir():
        raise ErrorArchivoJugadores("La carpeta del reporte no existe.")
    return ruta


def escribir_reporte_jugadores(informe, *, archivo, destino):
    """Exporta valores originales y resultados derivados, sin fórmulas activas."""
    libro = Workbook()
    resumen = libro.active
    resumen.title = "Resumen"
    resumen.append(["Dato", "Valor"])
    for nombre, valor in (
        ("Archivo procesado", str(archivo)), ("Temporada", informe.temporada),
        ("Modo", informe.modo), ("Fecha de referencia", informe.fecha_referencia),
        ("Rollback del conjunto seleccionado", informe.rollback),
        *informe.resumen.items(),
        ("Nota", "Existentes y conflictos pueden coincidir con revisión. Importables en modo importar indica filas seleccionadas antes de escribir."),
    ):
        resumen.append([nombre, valor])
    revision = libro.create_sheet("Revision")
    revision.append([
        "Fila física", "N°", "NOMBRES", "APELLIDOS", "RUT", "F.N", "CAT.",
        "Categoría calculada", "CONTACTO", "AVTIVO", "GENERO", "FEDERADO",
        "FEDERADO normalizado", "Estado del resultado", "Acción administrativa", "Motivos",
    ])
    apoderados = libro.create_sheet("Apoderados_pendientes")
    encabezados_administrativos = [
        "Fila física", "N°", "NOMBRES", "APELLIDOS", "RUT del jugador",
        "Categoría calculada", "CONTACTO original", "Estado de importación", "Situación administrativa",
    ]
    apoderados.append(encabezados_administrativos)
    contactos = None
    for fila in informe.filas:
        original = fila.originales
        requiere_revision = fila.estado in ("REVISIÓN", "POLÍTICA PENDIENTE")
        contacto_invalido = any(i.codigo == "CONTACTO_INVALIDO" for i in fila.incidencias)
        referencia = _referencia_categoria(fila.categoria_calculada)
        menor = referencia is not None and referencia[0] in ABREVIATURAS[:-1]
        datos_administrativos = [
            fila.fila, fila.numero, original["NOMBRES"], original["APELLIDOS"],
            original["RUT"], fila.categoria_calculada, original["CONTACTO"], fila.estado,
        ]
        if menor:
            if contacto_invalido:
                situacion = "Contacto informado con formato inválido; revisar y completar datos del apoderado."
            elif fila.normalizados.get("telefono"):
                situacion = "Teléfono del apoderado disponible; no persistido. Completar identidad, RUT, correo y parentesco del apoderado."
            else:
                situacion = "Sin contacto disponible. Completar contacto e identidad del apoderado."
            apoderados.append([*datos_administrativos, situacion])
        elif contacto_invalido and not requiere_revision:
            # Los contactos inválidos T/C siguen visibles sin mezclarse con errores bloqueantes.
            if contactos is None:
                contactos = libro.create_sheet("Contactos_pendientes")
                contactos.append(encabezados_administrativos)
            contactos.append([
                *datos_administrativos,
                "CONTACTO inválido; revisar el original con Administración. No persistido; no bloquea al jugador.",
            ])
        if requiere_revision:
            motivos = [i.mensaje for i in fila.incidencias if i.nivel != "INFO" or i.codigo == "CONTACTO_APODERADO"]
            revision.append([
                fila.fila, fila.numero, original["NOMBRES"], original["APELLIDOS"],
                original["RUT"], original["F.N"], original["CAT."], fila.categoria_calculada,
                original["CONTACTO"], original["AVTIVO"], original["GENERO"], original["FEDERADO"],
                fila.normalizados.get("federado"), fila.estado, "Corregir/revisar datos", "\n".join(motivos),
            ])
    for hoja in libro:
        hoja.freeze_panes = "A2"
        hoja.sheet_view.showGridLines = False
        for celda in hoja[1]:
            celda.font = Font(name="Calibri", bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="234E70")
        for fila in hoja.iter_rows(min_row=2):
            for celda in fila:
                if isinstance(celda.value, str):
                    # Incluye fórmulas originales y teléfonos que empiezan con +.
                    celda.data_type = "s"
                celda.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(celda.value, (date, datetime)):
                    celda.number_format = "dd-mm-yyyy"
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
    resumen.column_dimensions["A"].width = 40
    resumen.column_dimensions["B"].width = 110
    resumen.row_dimensions[resumen.max_row].height = 36
    revision.auto_filter.ref = revision.dimensions
    for columna, ancho in zip("ABCDEFGHIJKLMNOP", (12, 10, 24, 28, 18, 16, 12, 30, 24, 12, 12, 14, 22, 25, 48, 100)):
        revision.column_dimensions[columna].width = ancho
    for numero in range(2, revision.max_row + 1):
        motivos = revision.cell(numero, 16).value or ""
        lineas = sum(max(1, math.ceil(len(linea) / 90)) for linea in motivos.split("\n"))
        revision.row_dimensions[numero].height = max(48, 16 * lineas + 8)
    for hoja in (apoderados, contactos):
        if hoja is None:
            continue
        hoja.auto_filter.ref = hoja.dimensions
        for columna, ancho in zip("ABCDEFGHI", (12, 10, 24, 28, 18, 30, 24, 25, 100)):
            hoja.column_dimensions[columna].width = ancho
        for numero in range(2, hoja.max_row + 1):
            hoja.row_dimensions[numero].height = 48
    try:
        # Apertura exclusiva: ni siquiera un archivo aparecido durante la carga se sobrescribe.
        with Path(destino).open("xb") as salida:
            libro.save(salida)
    finally:
        libro.close()
