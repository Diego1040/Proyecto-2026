"""Libros generados en memoria; no dependen del XLSX manual del club."""

from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from django.apps import apps
from django.core import mail
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import connection
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900, to_excel

from .importacion_jugadores import (
    ENCABEZADOS, ErrorArchivoJugadores, ErrorImportacionJugadores, PoliticaImportacion,
    escribir_reporte_jugadores, importar_archivo_jugadores,
    normalizar_fecha_excel, validar_archivo_jugadores,
)
from .forms import InscripcionForm, InscripcionJugadorForm
from .models import (
    Apoderado, ApoderadoJugador, Categoria, HistorialCategoria, Jugador,
    ReglaCategoria, SolicitudInscripcion, Usuario,
)
from .services import asignar_categoria_automatica, cambiar_categoria_manual, crear_solicitud_inscripcion


REFERENCIA = date(2026, 10, 9)


def fila_base(**cambios):
    datos = {
        "N°": 1, "NOMBRES": " Alex ", "APELLIDOS": " Jugador ",
        "F.N": date(2010, 4, 14), "CAT.": "U17", "CONTACTO": "+56 9 1234 5678",
        "RUT": "11.111.111-1", "AVTIVO": "SI", "GENERO": "M", "FEDERADO": "NO",
    }
    datos.update(cambios)
    return datos


def libro_xlsx(filas, encabezados=ENCABEZADOS, hoja="Hoja1", epoch=CALENDAR_WINDOWS_1900):
    libro = Workbook()
    libro.epoch = epoch
    tabla = libro.active
    tabla.title = hoja
    tabla.append(list(encabezados))
    for fila in filas:
        tabla.append([fila.get(nombre.strip().upper()) for nombre in encabezados])
    archivo = BytesIO()
    libro.save(archivo)
    libro.close()
    archivo.seek(0)
    return archivo


class ImportacionJugadoresTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("cargar_categorias", temporada=2026, verbosity=0)

    def validar(self, filas, **opciones):
        return validar_archivo_jugadores(
            libro_xlsx(filas), temporada=opciones.pop("temporada", 2026),
            fecha_referencia=REFERENCIA, **opciones,
        )

    def codigos(self, fila):
        return {i.codigo for i in fila.incidencias}

    def crear_jugador(self, **cambios):
        datos = {
            "rut": "11111111-1", "nombres": "Alex", "apellidos": "Jugador",
            "fecha_nacimiento": date(2010, 4, 14), "rama": Jugador.Rama.VARONES,
            "estado": Jugador.Estado.ACTIVO,
            "categoria_actual": Categoria.objects.get(nombre="U17 Varones"),
        }
        datos.update(cambios)
        return Jugador.objects.create(**datos)

    def test_normaliza_sin_asumir_defaults_y_sin_persistir_contacto(self):
        resultado = self.validar([fila_base()]).filas[0]
        self.assertEqual(resultado.normalizados["rut"], "11111111-1")
        self.assertEqual(resultado.normalizados["telefono"], "+56912345678")
        self.assertEqual(resultado.normalizados["nombres"], "Alex")
        self.assertEqual(resultado.normalizados["rama"], "VARONES")
        self.assertEqual(resultado.normalizados["estado"], "ACTIVO")
        self.assertEqual(resultado.normalizados["procedencia"], "")
        self.assertEqual(resultado.categoria_calculada, "U17 Varones")
        self.assertEqual(resultado.estado, "OK")
        self.assertFalse(Jugador.objects.exists())

    def test_encabezados_con_espacios_minusculas_y_reordenados(self):
        encabezados = tuple(" " + c.lower() + " " for c in reversed(ENCABEZADOS))
        informe = validar_archivo_jugadores(libro_xlsx([fila_base()], encabezados), temporada=2026)
        self.assertEqual(informe.filas[0].normalizados["rut"], "11111111-1")

    def test_rechaza_encabezados_ausentes_duplicados_desconocidos_y_hoja(self):
        for encabezados in (
            ENCABEZADOS[:-1],
            ENCABEZADOS + (" nombres ",),
            ENCABEZADOS[:-1] + ("ACTIVO",),
        ):
            with self.subTest(encabezados=encabezados), self.assertRaises(ErrorArchivoJugadores):
                validar_archivo_jugadores(libro_xlsx([], encabezados), temporada=2026)
        with self.assertRaisesMessage(ErrorArchivoJugadores, "Hoja1"):
            validar_archivo_jugadores(libro_xlsx([], hoja="Otra"), temporada=2026)

    def test_archivo_corrupto_o_inexistente_se_reporta(self):
        for archivo in (BytesIO(b"no es XLSX"), "archivo_inexistente_importacion.xlsx"):
            with self.subTest(archivo=archivo), self.assertRaises(ErrorArchivoJugadores):
                validar_archivo_jugadores(archivo, temporada=2026)

    def test_temporada_es_explicita_y_valida(self):
        for temporada in (None, 0, -1, True, 10000):
            with self.subTest(temporada=temporada), self.assertRaises(ErrorArchivoJugadores):
                self.validar([], temporada=temporada)

    def test_fila_vacia_y_solo_correlativo_son_omitidas(self):
        informe = self.validar([{"N°": 87}, {}, fila_base(**{"N°": 42})])
        self.assertEqual([f.estado for f in informe.filas], ["OMITIDA", "OMITIDA", "OK"])
        self.assertEqual(informe.filas[2].fila, 4)
        self.assertEqual(informe.filas[2].numero, 42)
        self.assertEqual(informe.resumen["jugadores detectados"], 1)
        self.assertEqual(informe.resumen["omitidas"], 2)

    def test_todos_los_obligatorios_ausentes_se_reportan_en_una_fila(self):
        informe = self.validar([fila_base(**{c: None for c in ("NOMBRES", "APELLIDOS", "F.N", "RUT", "GENERO")})])
        fila = informe.filas[0]
        self.assertEqual(fila.estado, "REVISIÓN")
        self.assertTrue({"NOMBRES_AUSENTE", "APELLIDOS_AUSENTE", "FECHA_AUSENTE", "RUT_AUSENTE", "GENERO_INVALIDO"} <= self.codigos(fila))

    def test_full_clean_se_ejecuta_en_memoria_sin_excluir_campos_confirmados(self):
        original = Jugador.full_clean
        candidatos = []

        def verificar(candidato, **kwargs):
            candidatos.append(candidato)
            self.assertEqual(candidato.estado, "INACTIVO")
            self.assertEqual(candidato.procedencia, "")
            self.assertIsNone(candidato.pk)
            self.assertIsNone(candidato.usuario_id)
            self.assertEqual(kwargs["exclude"], [])
            return original(candidato, **kwargs)

        with patch.object(Jugador, "full_clean", verificar):
            self.validar([fila_base(**{"AVTIVO": None})])
        self.assertEqual(len(candidatos), 1)

    def test_longitud_y_tipo_de_nombres_se_validan(self):
        for valor in (123, "A" * 101):
            with self.subTest(valor=valor):
                fila = self.validar([fila_base(NOMBRES=valor)]).filas[0]
                self.assertEqual(fila.estado, "REVISIÓN")
                self.assertIn("NOMBRES_INVALIDO", self.codigos(fila))

    def test_fechas_reales_seriales_y_texto_con_anio_completo(self):
        nacimiento = date(2010, 4, 14)
        for valor in (nacimiento, datetime(2010, 4, 14), to_excel(datetime(2010, 4, 14)), "2010-04-14", "14/04/2010", "14-04-2010"):
            with self.subTest(valor=valor):
                fila = self.validar([fila_base(**{"F.N": valor})]).filas[0]
                self.assertEqual(fila.normalizados["fecha_nacimiento"], nacimiento)
                self.assertEqual(fila.categoria_calculada, "U17 Varones")

    def test_fecha_serial_utiliza_epoch_1904_del_libro(self):
        archivo = libro_xlsx([fila_base(**{"F.N": to_excel(datetime(2010, 4, 14), epoch=CALENDAR_MAC_1904)})], epoch=CALENDAR_MAC_1904)
        fila = validar_archivo_jugadores(archivo, temporada=2026).filas[0]
        self.assertEqual(fila.normalizados["fecha_nacimiento"], date(2010, 4, 14))

    def test_fechas_invalidas_ambiguas_futuras_y_horas_no_se_inventan(self):
        for valor in ("14/04/10", "31/02/2010", "texto", True, 0, -1, 60, datetime(2010, 4, 14, 12), date(2030, 1, 1)):
            with self.subTest(valor=valor):
                fila = self.validar([fila_base(**{"F.N": valor})]).filas[0]
                self.assertIn("FECHA_INVALIDA", self.codigos(fila))
                self.assertIsNone(fila.normalizados["fecha_nacimiento"])
        for valor in (float("nan"), float("inf")):
            with self.assertRaises(ValueError):
                normalizar_fecha_excel(valor, CALENDAR_WINDOWS_1900)

    def test_rut_invalido_no_se_corrige_y_numerico_entero_se_normaliza(self):
        fila = self.validar([fila_base(RUT="11111111-2")]).filas[0]
        self.assertIn("RUT_INVALIDO", self.codigos(fila))
        self.assertEqual(fila.originales["RUT"], "11111111-2")
        for valor in (111111111, 111111111.0, "011.111.111-1"):
            with self.subTest(valor=valor):
                self.assertEqual(self.validar([fila_base(RUT=valor)]).filas[0].normalizados["rut"], "11111111-1")
        for valor in (True, 111111111.5):
            self.assertIn("RUT_INVALIDO", self.codigos(self.validar([fila_base(RUT=valor)]).filas[0]))

    def test_duplicados_canonicos_reportan_todas_las_filas(self):
        informe = self.validar([fila_base(), fila_base(RUT="011111111-1", **{"N°": 8})])
        self.assertEqual(informe.resumen["conflictos"], 2)
        for fila in informe.filas:
            self.assertIn("RUT_DUPLICADO", self.codigos(fila))
            self.assertEqual(fila.estado, "REVISIÓN")

    def test_genero_f_m_y_valores_no_reconocidos(self):
        fila = self.validar([fila_base(GENERO=" f ")]).filas[0]
        self.assertEqual(fila.normalizados["rama"], "DAMAS")
        self.assertEqual(fila.categoria_calculada, "U17 Damas")
        for valor in ("X", "DAMAS", None):
            self.assertIn("GENERO_INVALIDO", self.codigos(self.validar([fila_base(GENERO=valor)]).filas[0]))

    def test_abreviaturas_categoria_y_edad_deportiva(self):
        for anio, referencia, esperado in (
            (2020, "U7", "U7 Mixto"), (2018, "U9", "U9 Mixto"),
            (2016, "U11", "U11 Mixto"), (2014, "U13", "U13 Varones"),
            (2012, "U15", "U15 Varones"), (2010, "U17", "U17 Varones"),
            (2008, "T/C", "Todo Competidor Varones"),
        ):
            with self.subTest(referencia=referencia):
                fila = self.validar([fila_base(**{"F.N": date(anio, 12, 31), "CAT.": referencia})]).filas[0]
                self.assertEqual(fila.categoria_calculada, esperado)
                self.assertNotIn("CATEGORIA_DISCREPANTE", self.codigos(fila))

    def test_categoria_discrepante_o_rama_diferente_no_crea_excepcion(self):
        for categoria in ("T/C", "U17 Damas"):
            fila = self.validar([fila_base(**{"CAT.": categoria})]).filas[0]
            self.assertIn("CATEGORIA_DISCREPANTE", self.codigos(fila))
        self.assertFalse(HistorialCategoria.objects.exists())

    def test_categoria_vacia_es_referencia_ausente_y_desconocida_es_revision(self):
        fila = self.validar([fila_base(**{"CAT.": None})]).filas[0]
        self.assertEqual(fila.categoria_calculada, "U17 Varones")
        self.assertIn("CATEGORIA_AUSENTE", self.codigos(fila))
        self.assertNotEqual(fila.estado, "REVISIÓN")
        fila = self.validar([fila_base(**{"CAT.": "Juvenil"})]).filas[0]
        self.assertIn("CATEGORIA_DESCONOCIDA", self.codigos(fila))

    def test_temporada_sin_regla_no_carga_categorias(self):
        anteriores = list(ReglaCategoria.objects.values())
        fila = self.validar([fila_base()], temporada=2027).filas[0]
        self.assertIn("SIN_REGLA_CATEGORIA", self.codigos(fila))
        self.assertEqual(list(ReglaCategoria.objects.values()), anteriores)

    def test_reglas_inactivas_o_superpuestas_se_reportan(self):
        ReglaCategoria.objects.filter(categoria__nombre="U17 Varones").update(activa=False)
        self.assertIn("SIN_REGLA_CATEGORIA", self.codigos(self.validar([fila_base()]).filas[0]))
        ReglaCategoria.objects.filter(categoria__nombre="U17 Varones").update(activa=True)
        ReglaCategoria.objects.filter(categoria__nombre="U15 Varones").update(edad_max=17)
        self.assertIn("REGLAS_CATEGORIA_INVALIDAS", self.codigos(self.validar([fila_base()]).filas[0]))

    def test_avtivo_si_no_vacio_interrogacion_y_otro(self):
        for valor, estado in (("SI", "ACTIVO"), (" no ", "INACTIVO"), (None, "INACTIVO")):
            fila = self.validar([fila_base(AVTIVO=valor)]).filas[0]
            self.assertEqual(fila.normalizados["estado"], estado)
            if valor is None:
                self.assertNotIn("ACTIVO_PENDIENTE", self.codigos(fila))
        for valor in ("?", "ACTIVO"):
            self.assertIn("ACTIVO_INVALIDO", self.codigos(self.validar([fila_base(AVTIVO=valor)]).filas[0]))

    def test_federado_triestado_se_conserva_y_valores_desconocidos_se_reportan(self):
        for valor, esperado in (("SI", True), ("NO", False), (None, False)):
            fila = self.validar([fila_base(FEDERADO=valor)]).filas[0]
            self.assertIs(fila.normalizados["federado"], esperado)
            self.assertIn("FEDERADO_LEIDO", self.codigos(fila))
        self.assertIn("FEDERADO_INVALIDO", self.codigos(self.validar([fila_base(FEDERADO="?")]).filas[0]))

    def test_contacto_opcional_y_sin_agregar_prefijos(self):
        fila = self.validar([fila_base(CONTACTO=None)]).filas[0]
        self.assertEqual(fila.normalizados["telefono"], "")
        self.assertNotEqual(fila.estado, "REVISIÓN")
        for contacto in ("912345678", "correo@example.com", 56912345678):
            fila = self.validar([fila_base(CONTACTO=contacto)]).filas[0]
            self.assertIn("CONTACTO_INVALIDO", self.codigos(fila))
            self.assertEqual(fila.estado, "OK")
            self.assertEqual(fila.normalizados["telefono"], "")

    def test_formulas_y_errores_excel_no_se_evaluan(self):
        for valor in ("=DATE(2010,4,14)", "#VALUE!"):
            fila = self.validar([fila_base(**{"F.N": valor})]).filas[0]
            self.assertIn("CELDA_NO_LITERAL", self.codigos(fila))
            self.assertIsNone(fila.normalizados["fecha_nacimiento"])

    def test_jugador_existente_coincidente_y_conflictivo_no_se_actualiza(self):
        jugador = self.crear_jugador()
        fila = self.validar([fila_base()]).filas[0]
        self.assertTrue(fila.existente)
        self.assertFalse(fila.conflicto)
        self.assertIn("JUGADOR_EXISTENTE", self.codigos(fila))
        fila = self.validar([fila_base(NOMBRES="Otra persona")]).filas[0]
        self.assertTrue(fila.conflicto)
        self.assertIn("JUGADOR_DIFERENTE", self.codigos(fila))
        jugador.refresh_from_db()
        self.assertEqual(jugador.nombres, "Alex")

    def test_cuenta_sin_jugador_se_reporta_sin_vincular(self):
        cuenta = Usuario.objects.create_user(rut="11111111-1")
        fila = self.validar([fila_base()]).filas[0]
        self.assertIn("RUT_EN_USUARIO", self.codigos(fila))
        self.assertTrue(fila.conflicto)
        self.assertFalse(Jugador.objects.filter(usuario=cuenta).exists())

    def test_rut_de_apoderado_se_reporta_y_respeta_validacion_identidad(self):
        Apoderado.objects.create(rut="11111111-1", nombres="Alex", apellidos="Tutor", telefono="+56912345678")
        fila = self.validar([fila_base()]).filas[0]
        self.assertIn("RUT_EN_APODERADO", self.codigos(fila))
        self.assertIn("VALIDACION_MODELO", self.codigos(fila))
        self.assertFalse(ApoderadoJugador.objects.exists())

    def test_coincidencia_canonica_detecta_rut_historico_no_normalizado(self):
        jugador = self.crear_jugador()
        Jugador.objects.filter(pk=jugador.pk).update(rut="011111111-1")
        self.assertTrue(self.validar([fila_base()]).filas[0].existente)

    def test_politicas_centralizadas_permiten_previsualizar_decisiones_confirmadas(self):
        politica = PoliticaImportacion(estado_vacio="ACTIVO", procedencia="INDEPENDIENTE")
        informe = self.validar([fila_base(AVTIVO=None)], politica=politica)
        self.assertEqual(informe.filas[0].estado, "OK")
        self.assertEqual(informe.resumen["OK"], 1)
        self.assertFalse(Jugador.objects.exists())

    def ejecutar_comando(self, archivo, **opciones):
        salida = StringIO()
        call_command("importar_jugadores", archivo, stdout=salida, **opciones)
        return salida.getvalue()

    def test_comando_sin_validar_y_sin_temporada_rechaza_la_operacion(self):
        with self.assertRaisesMessage(CommandError, "exactamente un modo"):
            self.ejecutar_comando("no_se_debe_abrir.xlsx", temporada=2026)
        with self.assertRaisesMessage(CommandError, "--temporada"):
            self.ejecutar_comando("no_se_debe_abrir.xlsx", validar=True)
        self.assertFalse(Jugador.objects.exists())

    def test_comando_muestra_fila_correlativo_categorias_incidencias_y_resumen(self):
        with TemporaryDirectory() as directorio:
            archivo = Path(directorio) / "jugadores.xlsx"
            archivo.write_bytes(libro_xlsx([fila_base(**{"N°": 42}), {"N°": 43}]).getvalue())
            salida = self.ejecutar_comando(str(archivo), validar=True, temporada=2026)
        for texto in ("Fila 2 | N° 42 | OK", "categoría Excel: U17", "categoría calculada: U17 Varones", "CONTACTO_APODERADO", "Fila 3 | N° 43 | OMITIDA", "filas examinadas: 2", "jugadores detectados: 1", "Escrituras realizadas: 0"):
            self.assertIn(texto, salida)

    def test_comando_informa_archivo_invalido_con_commanderror(self):
        with self.assertRaises(CommandError):
            self.ejecutar_comando("no_existe.xlsx", validar=True, temporada=2026)

    def test_validar_no_modifica_conteos_datos_historial_excepciones_ni_asignaciones(self):
        actor = Usuario.objects.create_user(rut="22222222-2")
        jugador = self.crear_jugador(categoria_actual=None)
        asignar_categoria_automatica(jugador, temporada=2026)
        cambiar_categoria_manual(jugador, Categoria.objects.get(nombre="U15 Varones"), actor, "Excepción autorizada existente")
        apoderado = Apoderado.objects.create(rut="33333333-3", nombres="Tutor", apellidos="Existente", telefono="+56987654321")
        ApoderadoJugador.objects.create(apoderado=apoderado, jugador=jugador, parentesco="Tutor", es_principal=True)
        SolicitudInscripcion.objects.create(jugador=jugador, procedencia="INDEPENDIENTE", consentimiento=True)
        modelos = [m for m in apps.get_models(include_auto_created=True) if m._meta.app_label in ("usuarios", "deportivo") and not m._meta.proxy]

        def snapshot():
            return {m._meta.label: {"conteo": m.objects.count(), "datos": list(m.objects.order_by("pk").values())} for m in modelos}

        antes = snapshot()
        escrituras = []

        def solo_select(execute, sql, params, many, context):
            if not sql.lstrip().upper().startswith("SELECT"):
                escrituras.append(sql)
                raise AssertionError("La previsualización intentó ejecutar SQL distinto de SELECT")
            return execute(sql, params, many, context)

        with TemporaryDirectory() as directorio:
            archivo = Path(directorio) / "solo_lectura.xlsx"
            archivo.write_bytes(libro_xlsx([
                fila_base(), fila_base(RUT="22222222-2"), fila_base(RUT="33333333-3"),
                fila_base(RUT="44444444-4"), fila_base(RUT="044444444-4"), {"N°": 99},
            ]).getvalue())
            with connection.execute_wrapper(solo_select):
                primera = self.ejecutar_comando(str(archivo), validar=True, temporada=2026)
                segunda = self.ejecutar_comando(str(archivo), validar=True, temporada=2026)
        self.assertEqual(escrituras, [])
        self.assertEqual(antes, snapshot())
        self.assertEqual(primera, segunda)
        self.assertIn("EXCEPCION_MANUAL", str(antes["usuarios.HistorialCategoria"]["datos"]))
        jugador.refresh_from_db()
        self.assertEqual(jugador.categoria_actual.nombre, "U15 Varones")

    def importar(self, filas):
        return importar_archivo_jugadores(libro_xlsx(filas), temporada=2026, fecha_referencia=REFERENCIA)

    def test_importa_estados_confirmados_y_excluye_interrogacion(self):
        informe = self.importar([
            fila_base(RUT="11111111-1", AVTIVO="SI"),
            fila_base(RUT="22222222-2", AVTIVO="NO"),
            fila_base(RUT="33333333-3", AVTIVO=None),
            fila_base(RUT="44444444-4", AVTIVO="?"),
        ])
        self.assertEqual(informe.resumen["creados"], 3)
        self.assertEqual(informe.resumen["revisión"], 1)
        self.assertEqual(dict(Jugador.objects.values_list("rut", "estado")), {
            "11111111-1": "ACTIVO", "22222222-2": "INACTIVO", "33333333-3": "INACTIVO",
        })
        self.assertEqual(HistorialCategoria.objects.count(), 3)

    def test_menor_sin_apoderado_se_importa_sin_telefono_ni_relaciones_ficticias(self):
        informe = self.importar([fila_base()])
        jugador = Jugador.objects.get()
        self.assertEqual(informe.resumen["creados"], 1)
        self.assertEqual(jugador.telefono, "")
        self.assertEqual(informe.filas[0].normalizados["telefono"], "+56912345678")
        self.assertEqual(informe.filas[0].normalizados["contacto_destino"], "APODERADO")
        for modelo in (Usuario, Apoderado, ApoderadoJugador, SolicitudInscripcion):
            self.assertEqual(modelo.objects.count(), 0)
        self.assertEqual(mail.outbox, [])
        self.assertIsNone(jugador.usuario_id)
        self.assertIsNone(jugador.fecha_ingreso)
        self.assertEqual(jugador.email, "")
        self.assertEqual(jugador.nombre_contacto_emergencia, "")
        self.assertEqual(jugador.telefono_contacto_emergencia, "")
        self.assertEqual(jugador.procedencia, "")
        self.assertEqual(jugador.club_anterior, "")
        jugador.full_clean()
        historial = jugador.historial_categorias.get()
        self.assertEqual(historial.tipo_cambio, "AUTOMATICO")
        self.assertIsNone(historial.categoria_anterior_id)
        self.assertIsNone(historial.cambiado_por_id)
        self.assertEqual(historial.categoria_nueva_id, jugador.categoria_actual_id)

    def test_contactos_de_todas_las_categorias_u_no_se_copian(self):
        filas = []
        for numero, (anio, categoria) in enumerate(((2020, "U7"), (2018, "U9"), (2016, "U11"), (2014, "U13"), (2012, "U15"), (2010, "U17")), 1):
            filas.append(fila_base(RUT=f"{str(numero) * 8}-{numero}", **{"F.N": date(anio, 1, 1), "CAT.": categoria}))
        self.assertEqual(self.importar(filas).resumen["creados"], 6)
        self.assertFalse(Jugador.objects.exclude(telefono="").exists())
        self.assertFalse(Apoderado.objects.exists())

    def test_contacto_tc_se_guarda_normalizado_y_vacio_se_permite(self):
        informe = self.importar([
            fila_base(**{"F.N": date(1990, 1, 1), "CAT.": "T/C"}),
            fila_base(RUT="22222222-2", CONTACTO=None, **{"F.N": date(1991, 1, 1), "CAT.": "T/C"}),
        ])
        self.assertEqual(informe.resumen["creados"], 2)
        self.assertEqual(Jugador.objects.get(rut="11111111-1").telefono, "+56912345678")
        self.assertEqual(Jugador.objects.get(rut="22222222-2").telefono, "")
        self.assertFalse(Jugador.objects.exclude(telefono_contacto_emergencia="").exists())

    def test_menores_con_contacto_invalido_se_importan_y_quedan_en_reporte(self):
        filas = []
        categorias = ((2020, "U7"), (2018, "U9"), (2016, "U11"), (2014, "U13"), (2012, "U15"), (2010, "U17"))
        contactos = (" 9 1234-5678 ", "correo@example.com", "=1+1", "#VALUE!", 56912345678, "--")
        for numero, ((anio, categoria), contacto) in enumerate(zip(categorias, contactos), 1):
            filas.append(fila_base(RUT=f"{str(numero) * 8}-{numero}", CONTACTO=contacto, **{"F.N": date(anio, 1, 1), "CAT.": categoria}))
        informe = self.importar(filas)
        self.assertEqual(informe.resumen["creados"], 6)
        self.assertEqual(informe.resumen["revisión"], 0)
        self.assertEqual(HistorialCategoria.objects.count(), 6)
        self.assertFalse(Jugador.objects.exclude(telefono="").exists())
        for modelo in (Usuario, Apoderado, ApoderadoJugador, SolicitudInscripcion):
            self.assertFalse(modelo.objects.exists())
        with TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "reporte.xlsx"
            escribir_reporte_jugadores(informe, archivo="entrada.xlsx", destino=ruta)
            libro = load_workbook(ruta, data_only=False)
            registros = list(libro["Apoderados_pendientes"].values)[1:]
            self.assertEqual(len(registros), 6)
            self.assertEqual(libro["Revision"].max_row, 1)
            for fila, registro, contacto in zip(informe.filas, registros, contactos):
                self.assertEqual(fila.estado, "CREADO")
                self.assertEqual(fila.normalizados["contacto_destino"], "APODERADO")
                incidencia = next(i for i in fila.incidencias if i.codigo == "CONTACTO_INVALIDO")
                self.assertEqual(incidencia.nivel, "ADMINISTRATIVA")
                self.assertEqual(registro[6], contacto)
                self.assertEqual(registro[7], "CREADO")
                self.assertEqual(registro[8], "Contacto informado con formato inválido; revisar y completar datos del apoderado.")
            self.assertEqual(libro["Apoderados_pendientes"]["G4"].data_type, "s")
            libro.close()

    def test_tc_con_contacto_invalido_se_importa_con_telefono_vacio_y_reporte(self):
        contacto = " +56 9 1234 567X "
        with TemporaryDirectory() as directorio:
            entrada, reporte = Path(directorio) / "entrada.xlsx", Path(directorio) / "reporte.xlsx"
            contenido = libro_xlsx([fila_base(CONTACTO=contacto, **{"F.N": date(1990, 1, 1), "CAT.": "T/C"})]).getvalue()
            entrada.write_bytes(contenido)
            salida = self.ejecutar_comando(str(entrada), importar=True, temporada=2026, reporte=str(reporte))
            jugador = Jugador.objects.get()
            self.assertEqual(jugador.telefono, "")
            self.assertEqual(jugador.categoria_actual.nombre, "Todo Competidor Varones")
            jugador.full_clean()
            self.assertIn("CONTACTO_INVALIDO", salida)
            self.assertIn("creados: 1", salida)
            self.assertIn("revisión: 0", salida)
            libro = load_workbook(reporte)
            registro = list(libro["Contactos_pendientes"].values)[1]
            self.assertEqual(registro[6], contacto)
            self.assertEqual(registro[7], "CREADO")
            self.assertIn("no bloquea", registro[8])
            self.assertIn("CONTACTO inválido", registro[8])
            self.assertEqual(libro["Revision"].max_row, 1)
            self.assertEqual(libro["Apoderados_pendientes"].max_row, 1)
            libro.close()
            self.assertEqual(entrada.read_bytes(), contenido)
            segunda = self.ejecutar_comando(str(entrada), importar=True, temporada=2026)
            self.assertIn("existentes sin cambios: 1", segunda)
        self.assertEqual(Jugador.objects.count(), 1)
        self.assertEqual(HistorialCategoria.objects.count(), 1)
        jugador.refresh_from_db()
        self.assertEqual(jugador.telefono, "")

    def test_contacto_invalido_conserva_original_sin_corregir_ni_inventar(self):
        for contacto in (" 9 1234-5678 ", "correo@example.com", 56912345678, "--", "=1+1", "#VALUE!"):
            with self.subTest(contacto=contacto), TemporaryDirectory() as directorio:
                entrada, reporte = Path(directorio) / "entrada.xlsx", Path(directorio) / "reporte.xlsx"
                contenido = libro_xlsx([fila_base(CONTACTO=contacto, **{"F.N": date(1990, 1, 1), "CAT.": "T/C"})]).getvalue()
                entrada.write_bytes(contenido)
                informe = validar_archivo_jugadores(entrada, temporada=2026, fecha_referencia=REFERENCIA)
                fila = informe.filas[0]
                self.assertEqual(fila.estado, "OK")
                self.assertEqual(fila.originales["CONTACTO"], contacto)
                self.assertEqual(fila.normalizados["telefono"], "")
                self.assertEqual(fila.normalizados["telefono_jugador"], "")
                self.assertIn("CONTACTO_INVALIDO", self.codigos(fila))
                escribir_reporte_jugadores(informe, archivo=entrada, destino=reporte)
                libro = load_workbook(reporte, data_only=False)
                self.assertEqual(libro["Contactos_pendientes"]["G2"].value, contacto)
                if isinstance(contacto, str):
                    self.assertEqual(libro["Contactos_pendientes"]["G2"].data_type, "s")
                libro.close()
                self.assertEqual(entrada.read_bytes(), contenido)
        self.assertFalse(Jugador.objects.exists())

    def test_titularidad_contacto_usa_categoria_calculada_con_cat_vacia(self):
        informe = self.importar([fila_base(**{"F.N": date(1990, 1, 1), "CAT.": None})])
        self.assertEqual(informe.resumen["creados"], 1)
        self.assertEqual(Jugador.objects.get().telefono, "+56912345678")

    def test_procedencia_default_normal_y_validacion_otro_club_se_conservan(self):
        jugador = Jugador(rut="11111111-1", nombres="Alex", apellidos="Jugador", fecha_nacimiento=date(1990, 1, 1), rama="VARONES")
        self.assertEqual(jugador.procedencia, "INDEPENDIENTE")
        jugador.full_clean()
        jugador.procedencia = ""
        jugador.full_clean()
        jugador.procedencia = "OTRO_CLUB"
        with self.assertRaises(ValidationError) as error:
            jugador.full_clean()
        self.assertIn("club_anterior", error.exception.message_dict)

    def test_formularios_y_servicio_de_inscripcion_siguen_exigiendo_procedencia(self):
        for clase in (InscripcionForm, InscripcionJugadorForm):
            formulario = clase({"procedencia": ""})
            self.assertTrue(formulario.fields["procedencia"].required)
            self.assertFalse(formulario.is_valid())
            self.assertIn("procedencia", formulario.errors)
        jugador = Jugador(rut="11111111-1", nombres="Alex", apellidos="Jugador", fecha_nacimiento=date(1990, 1, 1), rama="VARONES", procedencia="")
        with self.assertRaises(ValidationError) as error:
            crear_solicitud_inscripcion(jugador=jugador, consentimiento=True, temporada=2026)
        self.assertIn("procedencia", error.exception.message_dict)
        self.assertFalse(Jugador.objects.exists())
        self.assertFalse(SolicitudInscripcion.objects.exists())

    def test_cada_dato_critico_ausente_excluye_la_fila(self):
        for campo in ("NOMBRES", "APELLIDOS", "F.N", "RUT", "GENERO"):
            with self.subTest(campo=campo):
                informe = self.importar([fila_base(**{campo: None})])
                self.assertEqual(informe.resumen["creados"], 0)
                self.assertEqual(informe.filas[0].estado, "REVISIÓN")
        self.assertFalse(Jugador.objects.exists())

    def test_parcial_selecciona_validos_y_no_intenta_guardar_bloqueados(self):
        informe = self.importar([fila_base(), fila_base(RUT=None)])
        self.assertEqual(informe.resumen["importables"], 1)
        self.assertEqual(informe.resumen["creados"], 1)
        self.assertEqual(informe.resumen["revisión"], 1)
        self.assertEqual(Jugador.objects.count(), 1)

    def test_duplicados_del_archivo_no_se_importan_ni_se_elige_el_primero(self):
        informe = self.importar([fila_base(), fila_base(RUT="011111111-1")])
        self.assertEqual(informe.resumen["conflictos"], 2)
        self.assertEqual(informe.resumen["creados"], 0)

    def test_segunda_importacion_es_idempotente_en_jugadores_e_historiales(self):
        filas = [fila_base(), fila_base(RUT="22222222-2", **{"F.N": date(1990, 1, 1), "CAT.": "T/C"})]
        self.assertEqual(self.importar(filas).resumen["creados"], 2)
        jugadores = list(Jugador.objects.order_by("pk").values())
        historiales = list(HistorialCategoria.objects.order_by("pk").values())
        segunda = self.importar(filas)
        self.assertEqual(segunda.resumen["creados"], 0)
        self.assertEqual(segunda.resumen["existentes sin cambios"], 2)
        self.assertEqual(segunda.resumen["conflictos"], 0)
        self.assertEqual(jugadores, list(Jugador.objects.order_by("pk").values()))
        self.assertEqual(historiales, list(HistorialCategoria.objects.order_by("pk").values()))

    def test_existente_distinto_se_excluye_sin_sobrescribir(self):
        jugador = self.crear_jugador()
        antes = list(Jugador.objects.values())
        informe = self.importar([fila_base(NOMBRES="Otra identidad")])
        self.assertEqual(informe.resumen["creados"], 0)
        self.assertEqual(informe.resumen["conflictos"], 1)
        self.assertEqual(list(Jugador.objects.values()), antes)
        self.assertFalse(HistorialCategoria.objects.filter(jugador=jugador).exists())

    def test_existente_compatible_con_cuenta_ya_vinculada_se_omite(self):
        cuenta = Usuario.objects.create_user(rut="11111111-1")
        self.crear_jugador(usuario=cuenta)
        informe = self.importar([fila_base()])
        self.assertEqual(informe.filas[0].estado, "EXISTENTE SIN CAMBIOS")
        self.assertEqual(informe.resumen["creados"], 0)

    def test_contacto_adulto_diferente_se_reporta_sin_actualizar(self):
        self.importar([fila_base(**{"F.N": date(1990, 1, 1), "CAT.": "T/C"})])
        informe = self.importar([fila_base(CONTACTO="+56 9 8765 4321", **{"F.N": date(1990, 1, 1), "CAT.": "T/C"})])
        self.assertEqual(informe.resumen["conflictos"], 1)
        self.assertEqual(Jugador.objects.get().telefono, "+56912345678")

    def test_rollback_revierte_incluso_jugador_e_historial_anteriores_del_lote(self):
        llamadas = []

        def asignacion(jugador, **kwargs):
            llamadas.append(jugador.pk)
            if len(llamadas) == 2:
                raise RuntimeError("Falla simulada en segunda escritura")
            return asignar_categoria_automatica(jugador, **kwargs)

        with patch("usuarios.importacion_jugadores.asignar_categoria_automatica", side_effect=asignacion):
            with self.assertRaises(ErrorImportacionJugadores) as error:
                self.importar([fila_base(), fila_base(RUT="22222222-2")])
        self.assertEqual(len(llamadas), 2)
        self.assertTrue(error.exception.informe.rollback)
        self.assertEqual(error.exception.informe.resumen["creados"], 0)
        self.assertFalse(Jugador.objects.exists())
        self.assertFalse(HistorialCategoria.objects.exists())

    def test_fallo_full_clean_durante_escritura_revierte_lote(self):
        original = Jugador.full_clean

        def validar(jugador, **kwargs):
            if not kwargs and jugador.rut == "22222222-2":
                raise ValidationError("Fallo de validación al guardar")
            return original(jugador, **kwargs)

        with patch.object(Jugador, "full_clean", validar), self.assertRaises(ErrorImportacionJugadores):
            self.importar([fila_base(), fila_base(RUT="22222222-2")])
        self.assertFalse(Jugador.objects.exists())
        self.assertFalse(HistorialCategoria.objects.exists())

    def test_identidad_aparecida_despues_de_validar_revierte_seleccion(self):
        def validar(*args, **kwargs):
            informe = validar_archivo_jugadores(*args, **kwargs)
            Usuario.objects.create_user(rut="11111111-1")
            return informe

        with patch("usuarios.importacion_jugadores.validar_archivo_jugadores", side_effect=validar):
            with self.assertRaises(ErrorImportacionJugadores):
                self.importar([fila_base()])
        self.assertFalse(Jugador.objects.exists())
        self.assertFalse(HistorialCategoria.objects.exists())

    def test_fallo_validacion_del_historial_revierte_jugador(self):
        with patch.object(HistorialCategoria, "full_clean", side_effect=ValidationError("Historial inválido")):
            with self.assertRaises(ErrorImportacionJugadores):
                self.importar([fila_base()])
        self.assertFalse(Jugador.objects.exists())
        self.assertFalse(HistorialCategoria.objects.exists())

    def test_modos_son_excluyentes_tambien_desde_call_command(self):
        with self.assertRaises(CommandError):
            self.ejecutar_comando("no_se_debe_abrir.xlsx", validar=True, importar=True, temporada=2026)

    def test_reporte_validacion_contiene_originales_resumen_y_contacto_de_apoderado(self):
        informe = self.validar([fila_base(FEDERADO=None), fila_base(RUT=None, **{"N°": 2})])
        with TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "reporte.xlsx"
            escribir_reporte_jugadores(informe, archivo="entrada.xlsx", destino=ruta)
            libro = load_workbook(ruta)
            self.assertEqual(libro.sheetnames, ["Resumen", "Revision", "Apoderados_pendientes"])
            resumen = dict(libro["Resumen"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(resumen["importables"], 1)
            self.assertEqual(resumen["creados"], 0)
            self.assertEqual(libro["Revision"].max_row, 2)
            self.assertEqual(libro["Apoderados_pendientes"].max_row, 3)
            primera = list(libro["Apoderados_pendientes"].values)[1]
            self.assertEqual(primera[4], "11.111.111-1")
            self.assertEqual(primera[6], "+56 9 1234 5678")
            self.assertEqual(primera[7], "OK")
            self.assertIn("apoderado", primera[8])
            self.assertIs(libro["Revision"].cell(2, 13).value, False)
            self.assertIn("RUT ausente", libro["Revision"].cell(2, 16).value)
            libro.close()

    def _verificar_apoderados_pendientes(self, importar):
        contactos = ("+56 9 1234 5678", None, "teléfono desconocido") * 2
        categorias = ((2020, "U7"), (2018, "U9"), (2016, "U11"), (2014, "U13"), (2012, "U15"), (2010, "U17"))
        filas = [
            fila_base(RUT=f"{str(numero) * 8}-{numero}", CONTACTO=contacto, **{"N°": numero, "F.N": date(anio, 1, 1), "CAT.": categoria})
            for numero, ((anio, categoria), contacto) in enumerate(zip(categorias, contactos), 1)
        ]
        filas += [
            fila_base(RUT="77777777-7", **{"N°": 7, "F.N": date(1990, 1, 1), "CAT.": "T/C"}),
            fila_base(RUT=None, CONTACTO=None, **{"N°": 8}),
            {"N°": 9},
        ]
        informe = self.importar(filas) if importar else self.validar(filas)
        estado_esperado = "CREADO" if importar else "OK"
        self.assertEqual(informe.resumen["creados"], 7 if importar else 0)
        self.assertEqual(informe.resumen["revisión"], 1)
        estados_antes = [f.estado for f in informe.filas]
        modelos = [m for m in apps.get_models(include_auto_created=True) if m._meta.app_label in ("usuarios", "deportivo") and not m._meta.proxy]
        antes = {m._meta.label: list(m.objects.order_by("pk").values()) for m in modelos}
        with TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "reporte.xlsx"
            escribir_reporte_jugadores(informe, archivo="entrada.xlsx", destino=ruta)
            libro = load_workbook(ruta)
            self.assertEqual(libro.sheetnames, ["Resumen", "Revision", "Apoderados_pendientes"])
            hoja = libro["Apoderados_pendientes"]
            registros = list(hoja.values)[1:]
            self.assertEqual(len(registros), 7)
            situaciones = (
                "Teléfono del apoderado disponible; no persistido. Completar identidad, RUT, correo y parentesco del apoderado.",
                "Sin contacto disponible. Completar contacto e identidad del apoderado.",
                "Contacto informado con formato inválido; revisar y completar datos del apoderado.",
            ) * 2
            for numero, (registro, contacto, situacion, fila) in enumerate(zip(registros[:6], contactos, situaciones, informe.filas[:6]), 1):
                self.assertEqual(registro, (
                    numero + 1, numero, filas[numero - 1]["NOMBRES"], filas[numero - 1]["APELLIDOS"],
                    filas[numero - 1]["RUT"], fila.categoria_calculada, contacto, estado_esperado, situacion,
                ))
                self.assertEqual(fila.normalizados["telefono_jugador"], "")
            self.assertEqual(registros[-1][1], 8)
            self.assertEqual(registros[-1][7], "REVISIÓN")
            self.assertEqual(libro["Revision"].max_row, 2)
            self.assertEqual(libro["Revision"].cell(2, 2).value, 8)
            self.assertEqual(hoja.freeze_panes, "A2")
            self.assertEqual(hoja.auto_filter.ref, "A1:I8")
            libro.close()
        self.assertEqual([f.estado for f in informe.filas], estados_antes)
        self.assertEqual(antes, {m._meta.label: list(m.objects.order_by("pk").values()) for m in modelos})
        for modelo in (Apoderado, ApoderadoJugador):
            self.assertFalse(modelo.objects.exists())
        if importar:
            self.assertFalse(Jugador.objects.exclude(rut="77777777-7").exclude(telefono="").exists())
            self.assertEqual(Jugador.objects.get(rut="77777777-7").telefono, "+56912345678")

    def test_apoderados_pendientes_validacion_incluye_menores_con_y_sin_contacto(self):
        self._verificar_apoderados_pendientes(importar=False)

    def test_apoderados_pendientes_importacion_conserva_importabilidad_y_no_crea_apoderados(self):
        self._verificar_apoderados_pendientes(importar=True)

    def test_reporte_no_ejecuta_formulas_del_archivo_original(self):
        informe = self.validar([fila_base(NOMBRES="=1+1")])
        with TemporaryDirectory() as directorio:
            ruta = Path(directorio) / "reporte.xlsx"
            escribir_reporte_jugadores(informe, archivo="entrada.xlsx", destino=ruta)
            libro = load_workbook(ruta, data_only=False)
            self.assertEqual(libro["Revision"]["C2"].value, "=1+1")
            self.assertEqual(libro["Revision"]["C2"].data_type, "s")
            libro.close()

    def test_comando_validar_con_reporte_solo_ejecuta_select(self):
        def solo_select(execute, sql, params, many, context):
            self.assertTrue(sql.lstrip().upper().startswith("SELECT"), sql)
            return execute(sql, params, many, context)

        with TemporaryDirectory() as directorio:
            entrada, reporte = Path(directorio) / "entrada.xlsx", Path(directorio) / "reporte.xlsx"
            entrada.write_bytes(libro_xlsx([fila_base()]).getvalue())
            with connection.execute_wrapper(solo_select):
                self.ejecutar_comando(str(entrada), validar=True, temporada=2026, reporte=str(reporte))
            self.assertTrue(reporte.exists())
        self.assertFalse(Jugador.objects.exists())

    def test_comando_importar_con_reporte_y_repeticion(self):
        with TemporaryDirectory() as directorio:
            entrada = Path(directorio) / "entrada.xlsx"
            entrada.write_bytes(libro_xlsx([fila_base(), fila_base(RUT=None)]).getvalue())
            for vuelta in (1, 2):
                reporte = Path(directorio) / f"reporte_{vuelta}.xlsx"
                self.ejecutar_comando(str(entrada), importar=True, temporada=2026, reporte=str(reporte))
                libro = load_workbook(reporte)
                resumen = dict(libro["Resumen"].iter_rows(min_row=2, values_only=True))
                self.assertEqual(resumen["creados"], 1 if vuelta == 1 else 0)
                self.assertEqual(resumen["revisión"], 1)
                libro.close()
        self.assertEqual(Jugador.objects.count(), 1)
        self.assertEqual(HistorialCategoria.objects.count(), 1)

    def test_reporte_no_sobrescribe_entrada_ni_archivo_existente_antes_de_importar(self):
        with TemporaryDirectory() as directorio:
            entrada, reporte = Path(directorio) / "entrada.xlsx", Path(directorio) / "reporte.xlsx"
            contenido = libro_xlsx([fila_base()]).getvalue()
            entrada.write_bytes(contenido)
            reporte.write_bytes(b"conservar")
            for destino in (entrada, reporte):
                with self.assertRaises(CommandError):
                    self.ejecutar_comando(str(entrada), importar=True, temporada=2026, reporte=str(destino))
            self.assertEqual(entrada.read_bytes(), contenido)
            self.assertEqual(reporte.read_bytes(), b"conservar")
        self.assertFalse(Jugador.objects.exists())

    def test_reporte_de_rollback_no_anuncia_creaciones(self):
        with TemporaryDirectory() as directorio:
            entrada, reporte = Path(directorio) / "entrada.xlsx", Path(directorio) / "reporte.xlsx"
            entrada.write_bytes(libro_xlsx([fila_base()]).getvalue())
            with patch("usuarios.importacion_jugadores.asignar_categoria_automatica", side_effect=RuntimeError("Error simulado")):
                with self.assertRaisesMessage(CommandError, "rollback completo"):
                    self.ejecutar_comando(str(entrada), importar=True, temporada=2026, reporte=str(reporte))
            libro = load_workbook(reporte)
            resumen = dict(libro["Resumen"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(resumen["creados"], 0)
            self.assertIs(resumen["Rollback del conjunto seleccionado"], True)
            self.assertIn("rollback completo", libro["Revision"].cell(2, 16).value)
            libro.close()
        self.assertFalse(Jugador.objects.exists())
