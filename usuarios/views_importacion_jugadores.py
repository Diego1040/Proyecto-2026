"""Interfaz administrativa de HU19; todas las decisiones de jugadores son del motor."""

import hashlib
import logging
from datetime import date
from functools import wraps

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import FileResponse, Http404
from django.shortcuts import redirect, render
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_GET, require_http_methods, require_POST
from openpyxl import load_workbook

from deportivo.permisos import es_administracion

from . import temporales_importacion_jugadores as temporales
from .forms import ImportacionJugadoresForm
from .importacion_jugadores import (
    ErrorArchivoJugadores, ErrorImportacionJugadores,
    importar_archivo_jugadores, validar_archivo_jugadores,
)


logger = logging.getLogger(__name__)
CLAVE_SESION = "hu19_importacion"
PLANTILLA = "usuarios/importacion_jugadores.html"


def administracion_requerida(vista):
    @login_required
    @wraps(vista)
    def protegida(request, *args, **kwargs):
        if not request.user.is_active or not es_administracion(request.user):
            raise PermissionDenied
        return vista(request, *args, **kwargs)
    return protegida


def _sesion(request):
    return hashlib.sha256(request.session.session_key.encode("utf-8")).hexdigest()


def _proceso(request, proceso_id=None):
    proceso = request.session.get(CLAVE_SESION)
    if not proceso:
        if proceso_id is not None:
            raise Http404("Proceso no disponible.")
        return None
    if (proceso["usuario"] != request.user.pk or proceso["sesion"] != _sesion(request)
            or (proceso_id is not None and proceso["id"] != str(proceso_id))):
        raise Http404("Proceso no disponible.")
    if proceso["vence"] <= timezone.now().timestamp():
        temporales.eliminar(proceso["id"])
        request.session.pop(CLAVE_SESION)
        messages.error(request, "El proceso venció. Vuelve a seleccionar el Excel y validar.")
        return None
    return proceso


def _actualizar_informe(request, proceso, informe):
    # El reporte conserva la clasificación de errores y pendientes del motor.
    reporte = temporales.guardar_reporte(proceso["id"], informe, proceso["nombre"])
    proceso.update(
        reporte=reporte, resumen=informe.resumen,
        validacion=temporales.huella_validacion(informe), rollback=informe.rollback,
        vence=timezone.now().timestamp() + temporales.VIGENCIA_SEGUNDOS,
    )
    carpeta = temporales.directorio(proceso["id"])
    carpeta.touch()
    request.session[CLAVE_SESION] = proceso


def _tablas(proceso):
    """Proyecta las hojas existentes para HTML; no reclasifica filas ni contactos."""
    columnas = (
        ("Revision", "Errores y revisión", (
            "Fila física", "NOMBRES", "APELLIDOS", "RUT", "Categoría calculada", "Motivos",
        )),
        ("Apoderados_pendientes", "Apoderados pendientes", None),
        ("Contactos_pendientes", "Contactos pendientes", None),
    )
    tablas = []
    libro = load_workbook(
        temporales.ruta_reporte(proceso["id"], proceso["reporte"]), read_only=True, data_only=True,
    )
    try:
        for hoja, titulo, seleccion in columnas:
            if hoja not in libro.sheetnames:
                continue
            filas = libro[hoja].values
            encabezados = next(filas)
            indices = [encabezados.index(c) for c in seleccion] if seleccion else list(range(len(encabezados)))
            tablas.append({
                "titulo": titulo, "encabezados": [encabezados[i] for i in indices],
                "filas": [[fila[i] for i in indices] for fila in filas],
            })
    finally:
        libro.close()
    return tablas


def _pantalla(request, *, proceso=None, form=None, resultado=False):
    reporte_disponible = True
    try:
        tablas = _tablas(proceso) if proceso else []
    except FileNotFoundError:
        if not resultado:
            temporales.eliminar(proceso["id"])
            request.session.pop(CLAVE_SESION, None)
            messages.error(request, "La vista previa ya no está disponible. Vuelve a seleccionar el Excel y validar.")
            return redirect("importacion_jugadores")
        tablas = []
        reporte_disponible = False
    return render(request, PLANTILLA, {
        "form": form, "proceso": proceso, "tablas": tablas, "resultado": resultado,
        "reporte_disponible": reporte_disponible,
    })


@never_cache
@administracion_requerida
@require_http_methods(["GET", "POST"])
def importacion_jugadores(request):
    temporales.limpiar_vencidos()
    proceso = _proceso(request)
    if request.method == "GET":
        if proceso and proceso["estado"] != "validado":
            return redirect("importacion_jugadores_resultado", proceso_id=proceso["id"])
        return _pantalla(request, proceso=proceso, form=ImportacionJugadoresForm(
            initial={"temporada": timezone.localdate().year},
        ))

    form = ImportacionJugadoresForm(request.POST, request.FILES)
    if not form.is_valid():
        return _pantalla(request, form=form)
    nuevo_id = None
    try:
        nuevo_id = temporales.guardar(form.cleaned_data["archivo"])
        fecha = timezone.localdate()
        with (temporales.directorio(nuevo_id) / "entrada.xlsx").open("rb") as archivo:
            huella = temporales.huella_archivo(archivo)
            informe = validar_archivo_jugadores(
                archivo, temporada=form.cleaned_data["temporada"], fecha_referencia=fecha,
            )
        nuevo = {
            "id": nuevo_id, "usuario": request.user.pk, "sesion": _sesion(request),
            "temporada": informe.temporada, "fecha": fecha.isoformat(), "hash": huella,
            "nombre": form.cleaned_data["archivo"].name[:150], "estado": "validado",
            "reporte_final": False,
        }
        _actualizar_informe(request, nuevo, informe)
    except (ErrorArchivoJugadores, OSError, ValueError) as error:
        if nuevo_id:
            temporales.eliminar(nuevo_id)
        form.add_error("archivo", str(error) if isinstance(error, ErrorArchivoJugadores) else "No se pudo preparar el archivo y su reporte. Vuelve a intentarlo.")
        return _pantalla(request, form=form)
    if proceso:
        temporales.eliminar(proceso["id"])
    return redirect("importacion_jugadores")


@never_cache
@administracion_requerida
@require_POST
def importacion_jugadores_confirmar(request, proceso_id):
    proceso = _proceso(request, proceso_id)
    if not proceso:
        return redirect("importacion_jugadores")
    if proceso["estado"] != "validado":
        return redirect("importacion_jugadores_resultado", proceso_id=proceso_id)
    if request.POST.get("confirmar") != "si":
        messages.error(request, "Debes marcar la confirmación explícita para importar.")
        return redirect("importacion_jugadores")
    try:
        with temporales.confirmar_exclusivamente(proceso["id"]):
            return _confirmar(request, proceso)
    except FileExistsError:
        return render(request, PLANTILLA, {"en_curso": True}, status=409)
    except ErrorArchivoJugadores as error:
        temporales.eliminar(proceso["id"])
        request.session.pop(CLAVE_SESION, None)
        messages.error(request, str(error))
        return redirect("importacion_jugadores")
    except FileNotFoundError:
        request.session.pop(CLAVE_SESION, None)
        messages.error(request, "El Excel ya no está disponible. Vuelve a cargarlo y validar.")
        return redirect("importacion_jugadores")


def _confirmar(request, proceso):
    entrada = temporales.directorio(proceso["id"]) / "entrada.xlsx"
    with entrada.open("rb") as archivo:
        if temporales.huella_archivo(archivo) != proceso["hash"]:
            raise ErrorArchivoJugadores("El archivo cambió. Vuelve a seleccionarlo y validar.")
        opciones = {"temporada": proceso["temporada"], "fecha_referencia": date.fromisoformat(proceso["fecha"])}
        informe = validar_archivo_jugadores(archivo, **opciones)
        if temporales.huella_validacion(informe) != proceso["validacion"]:
            _actualizar_informe(request, proceso, informe)
            messages.warning(request, "La validación cambió desde la vista previa. Revisa nuevamente el resumen y el reporte antes de confirmar.")
            return redirect("importacion_jugadores")
        if not informe.resumen["importables"]:
            messages.error(request, "No hay jugadores nuevos válidos para importar.")
            return redirect("importacion_jugadores")
        archivo.seek(0)
        try:
            informe = importar_archivo_jugadores(archivo, **opciones)
        except ErrorImportacionJugadores as error:
            informe = error.informe

    proceso.update(
        estado="revertido" if informe.rollback else "importado", resumen=informe.resumen,
        rollback=informe.rollback,
    )
    try:
        _actualizar_informe(request, proceso, informe)
        proceso["reporte_final"] = True
    except (OSError, ValueError):
        logger.exception("No se pudo generar el reporte final HU19.")
        proceso["reporte_final"] = False
    finally:
        entrada.unlink(missing_ok=True)
        request.session[CLAVE_SESION] = proceso
    return redirect("importacion_jugadores_resultado", proceso_id=proceso["id"])


@never_cache
@administracion_requerida
@require_GET
def importacion_jugadores_reporte(request, proceso_id):
    proceso = _proceso(request, proceso_id)
    if not proceso:
        return redirect("importacion_jugadores")
    try:
        archivo = temporales.ruta_reporte(proceso["id"], proceso["reporte"]).open("rb")
    except FileNotFoundError:
        raise Http404("El reporte ya no está disponible.")
    return FileResponse(
        archivo, as_attachment=True, filename="reporte-jugadores.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@never_cache
@administracion_requerida
@require_GET
def importacion_jugadores_resultado(request, proceso_id):
    proceso = _proceso(request, proceso_id)
    if not proceso:
        return redirect("importacion_jugadores")
    if proceso["estado"] == "validado":
        return redirect("importacion_jugadores")
    temporales.limpiar_vencidos()
    return _pantalla(request, proceso=proceso, resultado=True, form=ImportacionJugadoresForm(
        initial={"temporada": timezone.localdate().year},
    ))
