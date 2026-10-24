"""Integración web HU19 con el motor real, Excel en memoria y temporales aislados."""

import json
import os
from datetime import date
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from uuid import uuid4

from django.apps import apps
from django.contrib.auth.models import Group
from django.core import mail
from django.core.exceptions import ImproperlyConfigured, PermissionDenied
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import Client, RequestFactory, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from . import temporales_importacion_jugadores as temporales
from .importacion_jugadores import importar_archivo_jugadores, validar_archivo_jugadores
from .models import (
    Apoderado, ApoderadoJugador, Categoria, HistorialCategoria, Jugador,
    ReglaCategoria, SolicitudInscripcion, Usuario,
)
from .services import asignar_categoria_automatica
from .test_importacion_jugadores import fila_base, libro_xlsx
from .views_importacion_jugadores import CLAVE_SESION, importacion_jugadores


REFERENCIA = date(2026, 10, 18)


class ImportacionJugadoresWebTest(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command("cargar_categorias", temporada=2026, verbosity=0)
        cls.administracion = Usuario.objects.create_user(rut="99999999-9")
        cls.grupo = Group.objects.create(name="Administración")
        cls.administracion.groups.add(cls.grupo)

    def setUp(self):
        temporal = TemporaryDirectory()
        self.addCleanup(temporal.cleanup)
        self.carpeta = Path(temporal.name)
        opciones = override_settings(HU19_TEMP_DIR=temporal.name)
        opciones.enable()
        self.addCleanup(opciones.disable)
        reloj = patch("usuarios.views_importacion_jugadores.timezone.localdate", return_value=REFERENCIA)
        reloj.start()
        self.addCleanup(reloj.stop)
        self.client.force_login(self.administracion)
        self.url = reverse("importacion_jugadores")

    def subir(self, filas=None, *, contenido=None, nombre="historicos.xlsx", temporada=2026, cliente=None):
        if contenido is None:
            contenido = libro_xlsx(filas if filas is not None else [fila_base()]).getvalue()
        return (cliente or self.client).post(self.url, {
            "temporada": temporada,
            "archivo": SimpleUploadedFile(nombre, contenido),
        })

    def proceso(self):
        return self.client.session[CLAVE_SESION]

    def ruta(self, etapa, identificador=None):
        return reverse("importacion_jugadores_" + etapa, args=[identificador or self.proceso()["id"]])

    def confirmar(self, **datos):
        return self.client.post(self.ruta("confirmar"), {"confirmar": "si", **datos})

    def snapshot(self, etiquetas=("usuarios", "deportivo", "finanzas")):
        return {
            modelo._meta.label: list(modelo.objects.order_by("pk").values())
            for modelo in apps.get_models(include_auto_created=True)
            if modelo._meta.app_label in etiquetas and not modelo._meta.proxy
        }

    def descargar(self):
        respuesta = self.client.get(self.ruta("reporte"))
        self.assertEqual(respuesta.status_code, 200)
        self.assertTrue(respuesta.streaming)
        self.assertIn("attachment", respuesta["Content-Disposition"])
        contenido = b"".join(respuesta.streaming_content)
        respuesta.close()
        return contenido

    def test_administracion_activa_sin_staff_accede_y_panel_muestra_enlace(self):
        respuesta = self.client.get(self.url)
        self.assertContains(respuesta, "Seleccionar archivo y temporada")
        self.assertContains(respuesta, 'enctype="multipart/form-data"')
        self.assertIn("no-store", respuesta["Cache-Control"])
        self.assertContains(self.client.get(reverse("panel")), self.url)
        self.assertFalse(self.administracion.is_staff)

    def test_anonimo_no_accede_a_ninguna_etapa(self):
        self.client.logout()
        for etapa in (None, "reporte", "confirmar", "resultado"):
            with self.subTest(etapa=etapa):
                url = self.url if etapa is None else self.ruta(etapa, uuid4())
                respuesta = self.client.post(url) if etapa == "confirmar" else self.client.get(url)
                self.assertEqual(respuesta.status_code, 302)
                self.assertIn(reverse("login"), respuesta.url)

    def test_solo_grupo_exacto_administracion_autoriza_todas_las_etapas(self):
        usuario = Usuario.objects.create_user(rut="88888888-8")
        for atributos, grupo in (
            ({}, None), ({"is_staff": True}, None),
            ({"is_staff": True, "is_superuser": True}, None),
            ({}, "Administracion"), ({}, "Entrenador"), ({}, "Tesorería"),
        ):
            with self.subTest(atributos=atributos, grupo=grupo):
                usuario.is_staff = atributos.get("is_staff", False)
                usuario.is_superuser = atributos.get("is_superuser", False)
                usuario.save()
                usuario.groups.clear()
                if grupo:
                    usuario.groups.add(Group.objects.get_or_create(name=grupo)[0])
                self.client.force_login(usuario)
                for etapa in (None, "reporte", "confirmar", "resultado"):
                    url = self.url if etapa is None else self.ruta(etapa, uuid4())
                    respuesta = self.client.post(url) if etapa == "confirmar" else self.client.get(url)
                    self.assertEqual(respuesta.status_code, 403)
                self.assertNotContains(self.client.get(reverse("panel")), self.url, status_code=302 if grupo == "Tesorería" else 200)
        self.assertFalse(Jugador.objects.exists())

    def test_administracion_inactiva_no_accede(self):
        self.administracion.is_active = False
        self.administracion.save()
        request = RequestFactory().get(self.url)
        request.user = self.administracion
        with self.assertRaises(PermissionDenied):
            importacion_jugadores(request)
        self.assertEqual(self.client.get(self.url).status_code, 302)

    def test_archivo_ausente_extension_corrupto_hoja_y_encabezados_invalidos(self):
        casos = (
            {"contenido": b"corrupto"},
            {"nombre": "historicos.txt"},
            {"contenido": libro_xlsx([], hoja="Otra").getvalue()},
            {"contenido": libro_xlsx([], encabezados=("NOMBRES",)).getvalue()},
            {"contenido": b""},
        )
        for opciones in casos:
            with self.subTest(opciones=tuple(opciones)):
                respuesta = self.subir(**opciones)
                self.assertEqual(respuesta.status_code, 200)
                self.assertTrue(respuesta.context["form"].errors)
                self.assertNotIn(CLAVE_SESION, self.client.session)
                self.assertEqual(list(self.carpeta.iterdir()), [])
        respuesta = self.client.post(self.url, {"temporada": 2026})
        self.assertIn("archivo", respuesta.context["form"].errors)
        self.assertFalse(Jugador.objects.exists())

    def test_limite_de_archivo_y_temporada_explicita(self):
        respuesta = self.subir(contenido=b"x" * (5 * 1024 * 1024 + 1))
        self.assertContains(respuesta, "no puede superar los 5 MB")
        for temporada in ("", "texto", 0, 10000):
            with self.subTest(temporada=temporada):
                respuesta = self.subir(temporada=temporada)
                self.assertIn("temporada", respuesta.context["form"].errors)
        self.assertNotIn(CLAVE_SESION, self.client.session)

    def test_validar_no_persiste_ninguna_entidad_de_negocio_y_sesion_es_pequena(self):
        antes = self.snapshot()
        self.assertRedirects(self.subir([fila_base(), fila_base(RUT=None)]), self.url)
        self.assertEqual(self.snapshot(), antes)
        proceso = self.proceso()
        self.assertEqual(proceso["resumen"]["importables"], 1)
        self.assertEqual(proceso["resumen"]["revisión"], 1)
        self.assertLess(len(json.dumps(proceso)), 2500)
        self.assertNotIn("+569", json.dumps(proceso))
        self.assertNotIn("11111111", json.dumps(proceso))
        self.assertNotIn("filas", proceso)
        carpeta = temporales.directorio(proceso["id"])
        self.assertFalse(carpeta.is_relative_to(Path.cwd()))
        self.assertTrue((carpeta / "entrada.xlsx").exists())
        with (carpeta / "entrada.xlsx").open("rb") as archivo:
            self.assertEqual(temporales.huella_archivo(archivo), proceso["hash"])

    def test_vista_previa_y_reporte_reutilizan_errores_y_pendientes_del_motor(self):
        self.subir([
            fila_base(CONTACTO="=1+1"), fila_base(RUT=None),
            fila_base(RUT="22222222-2", CONTACTO="inválido", **{"F.N": date(1990, 1, 1), "CAT.": "T/C"}),
        ])
        antes = self.snapshot()
        respuesta = self.client.get(self.url)
        for texto in ("Apoderados pendientes", "Contactos pendientes", "Errores y revisión", "RUT ausente", "=1+1"):
            self.assertContains(respuesta, texto)
        contenido = self.descargar()
        libro = load_workbook(BytesIO(contenido), data_only=False)
        self.assertEqual(libro.sheetnames, ["Resumen", "Revision", "Apoderados_pendientes", "Contactos_pendientes"])
        self.assertEqual(libro["Apoderados_pendientes"]["G2"].value, "=1+1")
        self.assertEqual(libro["Apoderados_pendientes"]["G2"].data_type, "s")
        self.assertEqual(libro["Resumen"]["B2"].value, "historicos.xlsx")
        libro.close()
        self.assertEqual(self.snapshot(), antes)

    def test_confirmar_importa_validos_sin_cuentas_apoderados_ni_efectos_en_otras_apps(self):
        filas = [fila_base(), fila_base(RUT=None), fila_base(
            RUT="22222222-2", **{"F.N": date(1990, 1, 1), "CAT.": "T/C"},
        )]
        self.subir(filas)
        entrada = temporales.directorio(self.proceso()["id"]) / "entrada.xlsx"
        otras_apps = self.snapshot(("deportivo", "finanzas"))
        cuentas = list(Usuario.objects.values())
        with patch("usuarios.views_importacion_jugadores.importar_archivo_jugadores", wraps=importar_archivo_jugadores) as importar:
            respuesta = self.confirmar(temporada=2027, categoria_id=999999, filas="[]", cantidad=0, archivo="otra-ruta.xlsx")
        self.assertRedirects(respuesta, self.ruta("resultado"))
        self.assertEqual(importar.call_args.kwargs, {"temporada": 2026, "fecha_referencia": REFERENCIA})
        self.assertEqual(Jugador.objects.count(), 2)
        self.assertEqual(HistorialCategoria.objects.count(), 2)
        self.assertFalse(entrada.exists())
        self.assertEqual(self.snapshot(("deportivo", "finanzas")), otras_apps)
        self.assertEqual(list(Usuario.objects.values()), cuentas)
        for modelo in (Apoderado, ApoderadoJugador, SolicitudInscripcion):
            self.assertFalse(modelo.objects.exists())
        self.assertFalse(Jugador.objects.exclude(usuario=None).exists())
        self.assertEqual(Jugador.objects.get(rut="11111111-1").telefono, "")
        self.assertEqual(Jugador.objects.get(rut="22222222-2").telefono, "+56912345678")
        self.assertEqual(mail.outbox, [])
        libro = load_workbook(BytesIO(self.descargar()))
        resumen = dict(libro["Resumen"].values)
        self.assertEqual(resumen["creados"], 2)
        self.assertEqual(resumen["Modo"], "importar")
        libro.close()

    def test_confirmacion_es_explicita_y_solo_post(self):
        self.subir()
        self.assertEqual(self.client.get(self.ruta("confirmar")).status_code, 405)
        self.assertRedirects(self.client.post(self.ruta("confirmar")), self.url)
        self.assertFalse(Jugador.objects.exists())
        self.assertEqual(self.client.post(self.ruta("reporte")).status_code, 405)
        self.assertEqual(self.client.post(self.ruta("resultado")).status_code, 405)

    def test_csrf_se_exige_al_validar_y_al_confirmar(self):
        cliente = Client(enforce_csrf_checks=True)
        cliente.force_login(self.administracion)
        cliente.get(self.url)
        self.assertEqual(self.subir(cliente=cliente).status_code, 403)
        archivo = SimpleUploadedFile("historicos.xlsx", libro_xlsx([fila_base()]).getvalue())
        cliente.post(self.url, {
            "archivo": archivo, "temporada": 2026,
            "csrfmiddlewaretoken": cliente.cookies["csrftoken"].value,
        })
        identificador = cliente.session[CLAVE_SESION]["id"]
        self.assertEqual(cliente.post(self.ruta("confirmar", identificador), {"confirmar": "si"}).status_code, 403)
        self.assertFalse(Jugador.objects.exists())

    def test_confirmar_sin_validar_o_con_otro_identificador_no_importa(self):
        self.assertEqual(self.client.post(self.ruta("confirmar", uuid4()), {"confirmar": "si"}).status_code, 404)
        self.subir()
        for etapa in ("confirmar", "reporte", "resultado"):
            url = self.ruta(etapa, uuid4())
            respuesta = self.client.post(url, {"confirmar": "si"}) if etapa == "confirmar" else self.client.get(url)
            self.assertEqual(respuesta.status_code, 404)
        self.assertFalse(Jugador.objects.exists())

    def test_proceso_esta_asociado_a_usuario_y_sesion(self):
        self.subir()
        otro_admin = Usuario.objects.create_user(rut="77777777-7")
        otro_admin.groups.add(self.grupo)
        for usuario in (otro_admin, self.administracion):
            cliente = Client()
            cliente.force_login(usuario)
            # Incluso copiar los metadatos no autoriza una sesión diferente.
            sesion = cliente.session
            sesion[CLAVE_SESION] = self.proceso()
            sesion.save()
            for etapa in ("confirmar", "reporte", "resultado"):
                url = self.ruta(etapa)
                respuesta = cliente.post(url, {"confirmar": "si"}) if etapa == "confirmar" else cliente.get(url)
                self.assertEqual(respuesta.status_code, 404)
        self.assertFalse(Jugador.objects.exists())

    def test_grupo_retirado_despues_de_validar_impide_confirmacion_y_descarga(self):
        self.subir()
        self.administracion.groups.clear()
        self.assertEqual(self.confirmar().status_code, 403)
        self.assertEqual(self.client.get(self.ruta("reporte")).status_code, 403)
        self.assertFalse(Jugador.objects.exists())

    def test_proceso_vencido_elimina_temporales_y_exige_nueva_validacion(self):
        self.subir()
        carpeta = temporales.directorio(self.proceso()["id"])
        url = self.ruta("confirmar")
        sesion = self.client.session
        proceso = sesion[CLAVE_SESION]
        proceso["vence"] = timezone.now().timestamp() - 1
        sesion[CLAVE_SESION] = proceso
        sesion.save()
        respuesta = self.client.post(url, {"confirmar": "si"}, follow=True)
        self.assertContains(respuesta, "El proceso venció")
        self.assertFalse(carpeta.exists())
        self.assertNotIn(CLAVE_SESION, self.client.session)
        self.assertFalse(Jugador.objects.exists())

    def test_limpieza_simple_elimina_huerfanos_vencidos_y_conserva_vigentes(self):
        self.subir()
        vigente = temporales.directorio(self.proceso()["id"])
        vencido = temporales.directorio(uuid4())
        vencido.mkdir()
        (vencido / "entrada.xlsx").write_bytes(b"temporal vencido")
        pasado = timezone.now().timestamp() - temporales.VIGENCIA_SEGUNDOS - 1
        os.utime(vencido, (pasado, pasado))
        ajeno = self.carpeta / "otra-carpeta"
        ajeno.mkdir()
        os.utime(ajeno, (pasado, pasado))
        self.client.get(self.url)
        self.assertFalse(vencido.exists())
        self.assertTrue(vigente.exists())
        self.assertTrue(ajeno.exists())

    def test_temporales_no_admiten_rutas_en_proyecto_ni_identificadores_de_ruta(self):
        with override_settings(HU19_TEMP_DIR=Path.cwd() / "temporales-hu19"):
            with self.assertRaises(ImproperlyConfigured):
                temporales.directorio(uuid4())
        with self.assertRaises(ValueError):
            temporales.directorio("../../entrada.xlsx")

    def test_archivo_alterado_no_se_importa_y_se_descarta(self):
        self.subir()
        carpeta = temporales.directorio(self.proceso()["id"])
        (carpeta / "entrada.xlsx").write_bytes(libro_xlsx([fila_base(NOMBRES="Alterado")]).getvalue())
        respuesta = self.confirmar()
        self.assertRedirects(respuesta, self.url)
        self.assertFalse(carpeta.exists())
        self.assertNotIn(CLAVE_SESION, self.client.session)
        self.assertFalse(Jugador.objects.exists())

    def test_temporal_ausente_se_informa_sin_importar(self):
        self.subir()
        (temporales.directorio(self.proceso()["id"]) / "entrada.xlsx").unlink()
        self.assertRedirects(self.confirmar(), self.url)
        self.assertFalse(Jugador.objects.exists())

    def test_reporte_perdido_exige_nueva_vista_previa_y_conserva_resultado_confirmado(self):
        self.subir()
        proceso = self.proceso()
        temporales.ruta_reporte(proceso["id"], proceso["reporte"]).unlink()
        respuesta = self.client.get(self.url, follow=True)
        self.assertContains(respuesta, "La vista previa ya no está disponible")
        self.assertNotIn(CLAVE_SESION, self.client.session)
        self.assertFalse(Jugador.objects.exists())
        self.subir()
        self.confirmar()
        proceso = self.proceso()
        temporales.ruta_reporte(proceso["id"], proceso["reporte"]).unlink()
        respuesta = self.client.get(self.ruta("resultado"))
        self.assertContains(respuesta, "Jugadores creados: 1")
        self.assertContains(respuesta, "El reporte ya no está disponible")
        self.assertRedirects(self.confirmar(), self.ruta("resultado"))
        self.assertEqual(Jugador.objects.count(), 1)

    def test_cambio_de_reglas_actualiza_previsualizacion_y_no_importa(self):
        self.subir()
        ReglaCategoria.objects.filter(categoria__nombre="U17 Varones").update(activa=False)
        with patch("usuarios.views_importacion_jugadores.importar_archivo_jugadores") as importar:
            respuesta = self.confirmar()
            importar.assert_not_called()
        self.assertRedirects(respuesta, self.url, fetch_redirect_response=False)
        self.assertContains(self.client.get(self.url), "La validación cambió")
        self.assertEqual(self.proceso()["resumen"]["importables"], 0)
        self.assertTrue((temporales.directorio(self.proceso()["id"]) / "entrada.xlsx").exists())
        self.assertFalse(Jugador.objects.exists())

    def test_detecta_cambio_relevante_aunque_resumen_y_categoria_id_sean_iguales(self):
        self.subir()
        anterior = self.proceso()
        Categoria.objects.filter(nombre="U17 Varones").update(nombre="U17 VARONES")
        self.assertRedirects(self.confirmar(), self.url)
        self.assertEqual(self.proceso()["resumen"], anterior["resumen"])
        self.assertNotEqual(self.proceso()["validacion"], anterior["validacion"])
        self.assertFalse(Jugador.objects.exists())
        self.assertRedirects(self.confirmar(), self.ruta("resultado"))
        self.assertEqual(Jugador.objects.count(), 1)

    def test_identidad_aparecida_despues_de_vista_previa_exige_revision(self):
        self.subir()
        Usuario.objects.create_user(rut="11111111-1")
        self.assertRedirects(self.confirmar(), self.url)
        self.assertEqual(self.proceso()["resumen"]["conflictos"], 1)
        self.assertFalse(Jugador.objects.exists())

    def test_refrescar_resultado_y_repetir_post_no_repite_importacion(self):
        self.subir()
        with patch("usuarios.views_importacion_jugadores.importar_archivo_jugadores", wraps=importar_archivo_jugadores) as importar:
            self.assertRedirects(self.confirmar(), self.ruta("resultado"))
            datos = self.snapshot()
            for _ in range(2):
                self.assertContains(self.client.get(self.ruta("resultado")), "Jugadores creados: 1")
                self.assertRedirects(self.confirmar(), self.ruta("resultado"))
            self.assertEqual(importar.call_count, 1)
        self.assertEqual(self.snapshot(), datos)

    def test_repetir_proceso_con_mismo_excel_no_duplica_jugadores_ni_historiales(self):
        self.subir()
        self.confirmar()
        antes = self.snapshot()
        carpeta_anterior = temporales.directorio(self.proceso()["id"])
        self.subir()
        self.assertFalse(carpeta_anterior.exists())
        self.assertEqual(self.proceso()["resumen"]["importables"], 0)
        self.assertEqual(self.proceso()["resumen"]["existentes sin cambios"], 1)
        self.assertRedirects(self.confirmar(), self.url)
        self.assertEqual(self.snapshot(), antes)

    def test_web_conserva_reglas_de_rut_duplicados_categoria_y_contacto(self):
        filas = [
            fila_base(RUT="11111111-2"),
            fila_base(RUT="22222222-2"), fila_base(RUT="022222222-2"),
            fila_base(RUT="33333333-3", **{"CAT.": "T/C"}),
            fila_base(RUT="44444444-4", CONTACTO="inválido"),
        ]
        esperado = validar_archivo_jugadores(libro_xlsx(filas), temporada=2026, fecha_referencia=REFERENCIA)
        self.subir(filas)
        self.assertEqual(self.proceso()["resumen"], esperado.resumen)
        self.assertEqual(self.proceso()["validacion"], temporales.huella_validacion(esperado))
        self.confirmar()
        self.assertEqual(list(Jugador.objects.values_list("rut", flat=True)), ["44444444-4"])

    def test_rollback_del_motor_se_muestra_y_no_deja_creaciones(self):
        self.subir([fila_base(), fila_base(RUT="22222222-2")])
        llamadas = []

        def asignar(jugador, **opciones):
            llamadas.append(jugador.pk)
            if len(llamadas) == 2:
                raise RuntimeError("Fallo simulado")
            return asignar_categoria_automatica(jugador, **opciones)

        with patch("usuarios.importacion_jugadores.asignar_categoria_automatica", side_effect=asignar):
            self.assertRedirects(self.confirmar(), self.ruta("resultado"))
        self.assertTrue(self.proceso()["rollback"])
        self.assertContains(self.client.get(self.ruta("resultado")), "se revirtió por completo")
        self.assertFalse(Jugador.objects.exists())
        self.assertFalse(HistorialCategoria.objects.exists())
        self.assertFalse((temporales.directorio(self.proceso()["id"]) / "entrada.xlsx").exists())

    def test_falla_reporte_final_no_oculta_importacion_ni_habilita_repeticion(self):
        self.subir()
        with self.assertLogs("usuarios.views_importacion_jugadores", level="ERROR"):
            with patch("usuarios.temporales_importacion_jugadores.escribir_reporte_jugadores", side_effect=OSError("Fallo simulado")):
                respuesta = self.confirmar()
        self.assertRedirects(respuesta, self.ruta("resultado"))
        self.assertEqual(self.proceso()["resumen"]["creados"], 1)
        self.assertFalse(self.proceso()["reporte_final"])
        self.assertFalse((temporales.directorio(self.proceso()["id"]) / "entrada.xlsx").exists())
        self.assertContains(self.client.get(self.ruta("resultado")), "No se pudo generar el reporte final")
        self.assertRedirects(self.confirmar(), self.ruta("resultado"))
        self.assertEqual(Jugador.objects.count(), 1)
        self.assertEqual(HistorialCategoria.objects.count(), 1)

    def test_doble_confirmacion_en_curso_no_llama_al_importador(self):
        self.subir()
        with temporales.confirmar_exclusivamente(self.proceso()["id"]):
            with patch("usuarios.views_importacion_jugadores.importar_archivo_jugadores") as importar:
                self.assertEqual(self.confirmar().status_code, 409)
                importar.assert_not_called()
        self.assertEqual(self.proceso()["estado"], "validado")
        self.assertFalse(Jugador.objects.exists())

    def test_menor_cronologico_en_tc_conserva_exactamente_comportamiento_actual(self):
        self.subir([fila_base(**{"F.N": date(2008, 12, 31), "CAT.": "T/C"})])
        self.confirmar()
        jugador = Jugador.objects.get()
        self.assertEqual(jugador.edad, 17)
        self.assertEqual(jugador.categoria_actual.nombre, "Todo Competidor Varones")
        self.assertEqual(jugador.telefono, "+56912345678")
        self.assertIsNone(jugador.usuario_id)
        self.assertFalse(ApoderadoJugador.objects.exists())
        libro = load_workbook(BytesIO(self.descargar()))
        self.assertEqual(libro["Apoderados_pendientes"].max_row, 1)
        libro.close()
